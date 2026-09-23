"""Read-only extraction of transaction-like rows from the hand-maintained workbook.

The original workbook is organised as side-by-side blocks of
``Date | Sources | Expenses`` columns, one block per school term, plus narrower
``Sources | Expenses`` blocks for Rent and Investments and a number of sidebar
areas holding totals, budgets and notes.  Nothing here ever opens the workbook for
writing: :func:`extract_candidates` only loads it twice (formulas and cached values)
so we can tell a real amount from a computed one.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
import re
import statistics

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from spending_tracker.parser import parse_line
from spending_tracker.schema import (
    PRECISION_EXACT,
    PRECISION_MONTH,
    PRECISION_PERIOD,
    PRECISION_UNKNOWN,
    TRANSACTION_COLUMNS,
)


STATUS_READY = "Ready"
STATUS_NEEDS_REVIEW = "Needs review"
STATUS_SKIPPED = "Skipped"

CONFIDENCE_HIGH = "High"
CONFIDENCE_MEDIUM = "Medium"
CONFIDENCE_LOW = "Low"

HEADER_SOURCES = {"sources", "source"}
HEADER_AMOUNT = {"expenses", "expense", "amount"}
HEADER_DATE = {"date", "dates"}
HEADER_SEARCH_ROWS = 6

TYPE_KEYWORDS = (
    ("income", "Income"),
    ("rent", "Rent"),
    ("investments", "Investment"),
    ("investment", "Investment"),
    ("expenses", "Expense"),
    ("expense", "Expense"),
)

# Labels that mark aggregates, budgets or headings rather than transactions.
AGGREGATE_PATTERN = re.compile(
    r"^(total|totals|subtotal|sub total|sum|average|avg|mean|balance|remaining|net|"
    r"grand total|estimate|estimated|actual|actual spent|budget|sources?|date|dates|"
    r"expenses?|income|investments?|w/ rent|left|difference)\b",
    re.IGNORECASE,
)
AGGREGATE_EXACT = {
    "total",
    "totals",
    "subtotal",
    "sum",
    "average",
    "avg",
    "balance",
    "net",
    "estimate",
    "actual spent",
    "expenses",
    "income",
    "sources",
    "date",
}


@dataclass(frozen=True)
class BlockSpec:
    """One ``Date | Sources | Expenses`` (or ``Sources | Expenses``) column block."""

    sheet: str
    header_row: int
    date_column: int | None
    description_column: int
    amount_column: int
    note_column: int
    transaction_type: str
    period: str
    label_cells: tuple[str, ...] = ()

    @property
    def columns(self) -> tuple[int, ...]:
        columns = [self.description_column, self.amount_column, self.note_column]
        if self.date_column is not None:
            columns.append(self.date_column)
        return tuple(sorted(columns))

    def describe(self) -> str:
        date_part = f"{get_column_letter(self.date_column)}/" if self.date_column else ""
        return (
            f"{self.sheet}: {date_part}{get_column_letter(self.description_column)}/"
            f"{get_column_letter(self.amount_column)} row {self.header_row}+ "
            f"[{self.transaction_type} · {self.period}]"
        )


@dataclass
class Candidate:
    """A migration candidate with its provenance and review status."""

    description: str
    amount: float | None
    transaction_type: str
    source_sheet: str
    source_reference: str
    source_line: str
    source_period: str
    date: date | None = None
    date_precision: str = PRECISION_UNKNOWN
    status: str = STATUS_READY
    confidence: str = CONFIDENCE_HIGH
    review_note: str = ""
    category: str = ""
    source_area: str = "Block"
    notes: list[str] = field(default_factory=list)

    def flag(self, status: str, note: str, confidence: str = CONFIDENCE_LOW) -> None:
        # Skipped beats Needs review; a row never becomes "more ready" after a flag.
        order = {STATUS_READY: 0, STATUS_NEEDS_REVIEW: 1, STATUS_SKIPPED: 2}
        if order[status] >= order[self.status]:
            self.status = status
        self.confidence = confidence
        self.notes.append(note)
        self.review_note = "; ".join(self.notes)

    def as_row(self) -> dict[str, object]:
        """Schema-shaped dict for spending_tracker.excel_store.append_transactions."""
        return {
            "Date": self.date,
            "Description": self.description,
            "Amount": self.amount,
            "Type": self.transaction_type,
            "Category": self.category,
            "Source Line": self.source_line,
            "Source Period": self.source_period,
            "Source Sheet": self.source_sheet,
            "Source Reference": self.source_reference,
            "Date Precision": self.date_precision,
            "Added At": "",
        }

    def as_report_row(self) -> dict[str, object]:
        row = self.as_row()
        row.pop("Added At")
        row["Source Area"] = self.source_area
        row["Status"] = self.status
        row["Confidence"] = self.confidence
        row["Review Note"] = self.review_note
        return row


def clean_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        # Excel's cached formula results carry float noise (63.77000000000001);
        # Source Line is read by a human, so show the money value.
        text = str(int(value)) if value.is_integer() else f"{round(value, 2):.2f}".rstrip("0").rstrip(".")
    else:
        text = str(value)
    return re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip()


def _normalized(value: object) -> str:
    return clean_text(value).lower()


def looks_like_aggregate(text: str) -> bool:
    normalized = _normalized(text)
    if not normalized:
        return False
    if normalized in AGGREGATE_EXACT:
        return True
    return bool(AGGREGATE_PATTERN.match(normalized))


def _type_from_label(label: str) -> str | None:
    normalized = _normalized(label)
    for keyword, transaction_type in TYPE_KEYWORDS:
        if keyword in normalized:
            return transaction_type
    return None


def _looks_numeric(text: str) -> bool:
    try:
        float(text.replace(",", "").replace("$", ""))
    except ValueError:
        return False
    return True


def find_blocks(ws) -> list[BlockSpec]:
    """Locate transaction blocks by their ``Sources`` / ``Expenses`` header pair."""
    blocks: list[BlockSpec] = []
    max_row = min(ws.max_row, HEADER_SEARCH_ROWS)
    for row in range(1, max_row + 1):
        for column in range(1, ws.max_column + 1):
            if _normalized(ws.cell(row, column).value) not in HEADER_SOURCES:
                continue
            if _normalized(ws.cell(row, column + 1).value) not in HEADER_AMOUNT:
                continue

            date_column = None
            if column > 1 and _normalized(ws.cell(row, column - 1).value) in HEADER_DATE:
                date_column = column - 1

            first_column = date_column or column
            labels = []
            for label_row in range(max(1, row - 2), row):
                for label_column in range(first_column, column + 2):
                    text = clean_text(ws.cell(label_row, label_column).value)
                    if text:
                        labels.append(text)

            transaction_type = "Expense"
            for label in labels:
                inferred = _type_from_label(label)
                if inferred is not None:
                    transaction_type = inferred
                    break
            remaining = [
                label
                for label in labels
                if _type_from_label(label) is None and not _looks_numeric(label)
            ]
            period = remaining[0] if remaining else ws.title

            blocks.append(
                BlockSpec(
                    sheet=ws.title,
                    header_row=row,
                    date_column=date_column,
                    description_column=column,
                    amount_column=column + 1,
                    note_column=column + 2,
                    transaction_type=transaction_type,
                    period=period,
                    label_cells=tuple(labels),
                )
            )
    return blocks


def inspect_workbook(workbook_path: Path) -> list[BlockSpec]:
    wb = load_workbook(workbook_path, read_only=False, data_only=True)
    try:
        return [block for ws in wb.worksheets for block in find_blocks(ws)]
    finally:
        wb.close()


MONTH_NAMES = {
    "january": 1, "jan": 1, "february": 2, "feb": 2, "march": 3, "mar": 3,
    "april": 4, "apr": 4, "may": 5, "june": 6, "jun": 6, "july": 7, "jul": 7,
    "august": 8, "aug": 8, "september": 9, "sep": 9, "sept": 9,
    "october": 10, "oct": 10, "november": 11, "nov": 11, "december": 12, "dec": 12,
}
MONTH_TOKEN_PATTERN = re.compile(
    r"(?<![a-z])(" + "|".join(sorted(MONTH_NAMES, key=len, reverse=True)) + r")\.?(?![a-z])",
    re.IGNORECASE,
)


def month_from_description(text: str) -> int | None:
    """Month number when a description names exactly one month, else None.

    The old workbook labels undated rent add-ons and utility bills by month
    ("Rent+ Aug.", "May Utils"), which is real month-level information.
    """
    found = {MONTH_NAMES[match.group(1).lower()] for match in MONTH_TOKEN_PATTERN.finditer(text)}
    if len(found) != 1:
        return None
    return found.pop()


def month_date_near_block(month: int, block_dates: list[date]) -> date | None:
    """First of ``month``, with the year closest to the dates the block actually has."""
    if not block_dates:
        return None
    start, end = min(block_dates), max(block_dates)
    best: date | None = None
    best_distance = None
    for year in range(start.year - 1, end.year + 2):
        try:
            candidate = date(year, month, 1)
        except ValueError:  # pragma: no cover - month is always 1..12 here
            continue
        if start <= candidate <= end:
            distance = 0
        else:
            distance = min(abs((candidate - start).days), abs((candidate - end).days))
        if best_distance is None or distance < best_distance:
            best, best_distance = candidate, distance
    return best


def _block_representative_date(dates: list[date]) -> date | None:
    """Midpoint of the exact dates seen inside a block.

    Used only as a sortable stand-in for rows whose day was never recorded; the row
    keeps Date Precision = Period so the uncertainty stays visible.
    """
    if not dates:
        return None
    ordinals = sorted(value.toordinal() for value in dates)
    return date.fromordinal(int(statistics.median_low(ordinals)))


def _collect_block_dates(ws_values, block: BlockSpec) -> list[date]:
    column = block.date_column or block.description_column
    dates = []
    for row in range(block.header_row + 1, ws_values.max_row + 1):
        value = ws_values.cell(row, column).value
        if isinstance(value, datetime):
            dates.append(value.date())
        elif isinstance(value, date):
            dates.append(value)
    return dates


def _extract_block(
    ws_values, ws_formulas, block: BlockSpec, sheet_representative: date | None = None
) -> tuple[list[Candidate], int]:
    block_dates = _collect_block_dates(ws_values, block)
    representative = _block_representative_date(block_dates)
    candidates: list[Candidate] = []
    blank_rows = 0

    for row in range(block.header_row + 1, ws_values.max_row + 1):
        raw_description = ws_values.cell(row, block.description_column).value
        raw_amount = ws_values.cell(row, block.amount_column).value
        amount_formula = ws_formulas.cell(row, block.amount_column).value
        raw_date = ws_values.cell(row, block.date_column).value if block.date_column else None
        raw_note = ws_values.cell(row, block.note_column).value

        if raw_description is None and raw_amount is None:
            blank_rows += 1
            continue

        date_text = clean_text(raw_date)
        description_text = clean_text(raw_description)
        note_text = clean_text(raw_note)
        source_line = " | ".join(part for part in (date_text, description_text, clean_text(raw_amount)) if part)

        reference = f"{get_column_letter(block.description_column)}{row}"
        candidate = Candidate(
            description=description_text,
            amount=None,
            transaction_type=block.transaction_type,
            source_sheet=block.sheet,
            source_reference=reference,
            source_line=source_line,
            source_period=block.period,
        )

        # Resolve the date.  Two-column blocks keep the date in the "Sources" column.
        date_value: date | None = None
        if block.date_column is not None:
            date_value = _as_date(raw_date)
        else:
            date_value = _as_date(raw_description)
            if date_value is not None:
                candidate.description = block.period if block.transaction_type == "Expense" else block.transaction_type
                if note_text:
                    candidate.description = f"{candidate.description} ({note_text.strip('()')})"
                candidate.source_line = " | ".join(
                    part for part in (date_text or description_text, clean_text(raw_amount)) if part
                )

        month_date = None
        if date_value is None:
            month = month_from_description(candidate.description)
            if month is not None:
                month_date = month_date_near_block(month, block_dates)

        if date_value is not None:
            candidate.date = date_value
            candidate.date_precision = PRECISION_EXACT
        elif month_date is not None:
            candidate.date = month_date
            candidate.date_precision = PRECISION_MONTH
            candidate.notes.append(
                f"Month taken from the description; day unknown, using {month_date.isoformat()}"
            )
            candidate.review_note = "; ".join(candidate.notes)
        elif representative is not None:
            candidate.date = representative
            candidate.date_precision = PRECISION_PERIOD
            candidate.notes.append(
                f"No transaction date in source; using block midpoint {representative.isoformat()} "
                f"for period '{block.period}'"
            )
            candidate.review_note = "; ".join(candidate.notes)
        elif sheet_representative is not None:
            # Some blocks (Investments) carry no dates at all, but the sheet they live
            # on does.  The transaction itself is not in doubt - only the day - so the
            # row keeps Period precision and says which midpoint stood in for the date.
            candidate.date = sheet_representative
            candidate.date_precision = PRECISION_PERIOD
            candidate.notes.append(
                f"No date anywhere in block '{block.period}'; using the midpoint of dated rows on "
                f"sheet '{block.sheet}' ({sheet_representative.isoformat()})"
            )
            candidate.review_note = "; ".join(candidate.notes)
        else:
            candidate.date = None
            candidate.date_precision = PRECISION_UNKNOWN
            candidate.flag(
                STATUS_NEEDS_REVIEW,
                f"No date information anywhere on sheet '{block.sheet}'",
                CONFIDENCE_LOW,
            )

        if not candidate.description:
            candidate.flag(STATUS_SKIPPED, "No description in row")
            candidates.append(candidate)
            continue

        if looks_like_aggregate(candidate.description):
            candidate.flag(STATUS_SKIPPED, "Description looks like a total, heading or label")
            candidates.append(candidate)
            continue

        # Formula amounts are never auto-migrated: they are either combined entries
        # or references to a total.  Checked before the value, because a workbook that
        # has not been recalculated has no cached value at all.
        if isinstance(amount_formula, str) and amount_formula.startswith("="):
            if isinstance(raw_amount, (int, float)) and not isinstance(raw_amount, bool):
                candidate.amount = round(abs(float(raw_amount)), 2)
            candidate.flag(
                STATUS_NEEDS_REVIEW,
                f"Amount cell is a formula ({amount_formula}) - may be a combined entry or a reference to a total",
            )
            candidates.append(candidate)
            continue

        if raw_amount is None:
            candidate.flag(STATUS_SKIPPED, "No amount in row")
            candidates.append(candidate)
            continue

        if not isinstance(raw_amount, (int, float)) or isinstance(raw_amount, bool):
            candidate.flag(STATUS_SKIPPED, f"Amount is not numeric ({clean_text(raw_amount)!r})")
            candidates.append(candidate)
            continue

        amount = float(raw_amount)
        if amount == 0:
            candidate.flag(STATUS_SKIPPED, "Amount is zero")
            candidates.append(candidate)
            continue

        if amount < 0:
            candidate.amount = round(abs(amount), 2)
            candidate.transaction_type = "Expense" if block.transaction_type == "Income" else "Income"
            candidate.flag(
                STATUS_NEEDS_REVIEW,
                f"Negative amount in a '{block.transaction_type}' column; read as a credit and "
                f"re-typed {candidate.transaction_type} - confirm",
            )
            candidates.append(candidate)
            continue

        candidate.amount = round(amount, 2)
        if note_text:
            candidate.notes.append(f"Source note: {note_text}")
            candidate.review_note = "; ".join(candidate.notes)
        candidates.append(candidate)

    return candidates, blank_rows


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _owned_columns(blocks: list[BlockSpec]) -> set[int]:
    owned: set[int] = set()
    for block in blocks:
        owned.update(block.columns)
    return owned


PERIOD_WORD_PATTERN = re.compile(
    r"^(fall|winter|spring|summer|autumn|q[1-4]|quarter|semester|term|january|february|march|"
    r"april|may|june|july|august|september|october|november|december)$",
    re.IGNORECASE,
)
QUANTITY_SUFFIX_PATTERN = re.compile(r"\*\s*\d+\s*$")


def _nearest_label(ws_values, row: int, column: int, owned: set[int]) -> str:
    for offset in range(1, 4):
        target = column - offset
        if target < 1 or target in owned:
            break
        text = clean_text(ws_values.cell(row, target).value)
        if text and not isinstance(ws_values.cell(row, target).value, (int, float)):
            return text
    for offset in range(1, 13):
        if row - offset < 1:
            break
        value = ws_values.cell(row - offset, column).value
        text = clean_text(value)
        if text and not isinstance(value, (int, float, datetime, date)):
            return text
    return ""


def _extract_sidebar(ws_values, ws_formulas, blocks: list[BlockSpec]) -> list[Candidate]:
    """Capture values that sit outside any recognised block.

    These are never auto-migrated: the sidebars mix genuine one-off entries
    (deposits, brokerage transfers) with totals and budget planning, so every hit is
    reported as "Needs review" with its cell reference.
    """
    owned = _owned_columns(blocks)
    candidates: list[Candidate] = []

    for row in range(1, ws_values.max_row + 1):
        for column in range(1, ws_values.max_column + 1):
            if column in owned:
                continue
            value = ws_values.cell(row, column).value
            formula = ws_formulas.cell(row, column).value
            if value is None:
                continue
            if isinstance(formula, str) and formula.startswith("="):
                continue

            reference = f"{get_column_letter(column)}{row}"
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                if value == 0:
                    continue
                label = _nearest_label(ws_values, row, column, owned)
                if not label or looks_like_aggregate(label):
                    continue
                transaction_type = _type_from_label(label) or "Expense"
                if value < 0 and transaction_type in {"Expense", "Rent"}:
                    transaction_type = "Income"
                candidate = Candidate(
                    description=label,
                    amount=round(abs(float(value)), 2),
                    transaction_type=transaction_type,
                    source_sheet=ws_values.title,
                    source_reference=reference,
                    source_line=f"{label} | {clean_text(value)}",
                    source_period=ws_values.title,
                    date=None,
                    date_precision=PRECISION_UNKNOWN,
                )
                candidate.source_area = "Sidebar"
                candidate.flag(
                    STATUS_NEEDS_REVIEW,
                    f"Value outside any transaction block (nearest label {label!r}); no date available",
                )
                candidates.append(candidate)
                continue

            text = clean_text(value)
            if not text or looks_like_aggregate(text):
                continue
            if QUANTITY_SUFFIX_PATTERN.search(text):
                continue
            parsed = parse_line(text, date.today())
            if parsed.status != "Ready" or parsed.amount is None:
                continue
            if PERIOD_WORD_PATTERN.match(parsed.description):
                continue
            candidate = Candidate(
                description=parsed.description,
                amount=round(float(parsed.amount), 2),
                transaction_type=parsed.type,
                source_sheet=ws_values.title,
                source_reference=reference,
                source_line=text,
                source_period=ws_values.title,
                date=None,
                date_precision=PRECISION_UNKNOWN,
            )
            candidate.source_area = "Sidebar"
            candidate.flag(
                STATUS_NEEDS_REVIEW,
                "Free-text note outside any transaction block; no date available",
            )
            candidates.append(candidate)

    return candidates


@dataclass
class Extraction:
    """Everything the read-only pass over the historical workbook found."""

    candidates: list[Candidate]
    blocks: list[BlockSpec]
    blank_rows_ignored: int
    source_workbook: str


def extract(workbook_path: Path, include_sidebar: bool = True) -> Extraction:
    """Extract every transaction-like row from the historical workbook (read-only)."""
    values_wb = load_workbook(workbook_path, data_only=True)
    formulas_wb = load_workbook(workbook_path, data_only=False)
    try:
        candidates: list[Candidate] = []
        all_blocks: list[BlockSpec] = []
        blank_rows = 0
        for ws_values in values_wb.worksheets:
            ws_formulas = formulas_wb[ws_values.title]
            blocks = find_blocks(ws_values)
            all_blocks.extend(blocks)
            sheet_dates = [
                value for block in blocks for value in _collect_block_dates(ws_values, block)
            ]
            sheet_representative = _block_representative_date(sheet_dates)
            for block in blocks:
                block_candidates, block_blanks = _extract_block(
                    ws_values, ws_formulas, block, sheet_representative
                )
                candidates.extend(block_candidates)
                blank_rows += block_blanks
            if include_sidebar:
                candidates.extend(_extract_sidebar(ws_values, ws_formulas, blocks))
        return Extraction(
            candidates=candidates,
            blocks=all_blocks,
            blank_rows_ignored=blank_rows,
            source_workbook=str(workbook_path),
        )
    finally:
        values_wb.close()
        formulas_wb.close()


def extract_candidates(workbook_path: Path, include_sidebar: bool = True) -> list[Candidate]:
    return extract(workbook_path, include_sidebar=include_sidebar).candidates


def ready_candidates(candidates: list[Candidate]) -> list[Candidate]:
    return [
        candidate
        for candidate in candidates
        if candidate.status == STATUS_READY and candidate.confidence == CONFIDENCE_HIGH and candidate.amount
    ]


__all__ = [
    "CONFIDENCE_HIGH",
    "CONFIDENCE_LOW",
    "CONFIDENCE_MEDIUM",
    "STATUS_NEEDS_REVIEW",
    "STATUS_READY",
    "STATUS_SKIPPED",
    "BlockSpec",
    "Candidate",
    "Extraction",
    "extract",
    "clean_text",
    "extract_candidates",
    "find_blocks",
    "inspect_workbook",
    "looks_like_aggregate",
    "month_date_near_block",
    "month_from_description",
    "ready_candidates",
    "TRANSACTION_COLUMNS",
]
