"""One-time Historical Review workflow for rows the migration held back.

The migration and extraction rules are not touched here.  This module reads the
``Needs review`` candidates the migration already produces, remembers what you decided
about each one in a small JSON file, and - only when you approve an item - appends a
normalized transaction to the tracker workbook with its provenance intact.

Nothing is ever approved automatically.  Suggestions are prefilled where the source is
unambiguous (a negative amount in a Rent column is a reimbursement, a sidebar entry
named after a brokerage is an investment); everything else is shown with low
confidence and waits for a human.
"""

from __future__ import annotations

from collections.abc import Mapping
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
import json
import re
import statistics

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from spending_tracker.categories import categorize
from spending_tracker.excel_store import append_transactions, backup_workbook, load_transactions
from spending_tracker.historical import Candidate, find_blocks
from spending_tracker.manage import record_audit
from spending_tracker.migration import dry_run, provenance_identity
from spending_tracker.paths import DEFAULT_DATA_DIR
from spending_tracker.schema import (
    PRECISION_EXACT,
    PRECISION_PERIOD,
    PRECISION_UNKNOWN,
    VALID_TYPES,
)


DEFAULT_REVIEW_STATE_PATH = DEFAULT_DATA_DIR / "historical_review.json"
STATE_VERSION = 1

STATUS_UNRESOLVED = "Unresolved"
STATUS_APPROVED = "Approved"
STATUS_IGNORED = "Ignored"

CONFIDENCE_HIGH = "High"
CONFIDENCE_LOW = "Low"

ACTION_APPROVE = "approve"
ACTION_APPROVE_EDITED = "approve-edited"
ACTION_IGNORE = "ignore"

AUDIT_APPROVE = "historical-review-approve"
AUDIT_IGNORE = "historical-review-ignore"

# A negative amount in a Rent column is money coming back out of rent.
_RENT_CREDIT_NOTE = "'rent' column"
# Words that make a credit in an expense column unambiguous income.
_CREDIT_WORDS = re.compile(r"\b(pay|payout|credit|cashback|cash back|refund|reimbursement|rebate)\b", re.I)
_INVESTMENT_WORDS = re.compile(r"\b(ibkr|robinhood|brokerage|vanguard|fidelity|schwab|investments?)\b", re.I)
_CELL_REFERENCE = re.compile(r"(?<![A-Z0-9_!$])(\$?[A-Z]{1,3}\$?\d{1,7})(?![\d(])")


# --- state -----------------------------------------------------------------------


def load_state(path: Path | None = None) -> dict[str, dict[str, object]]:
    """Read the decision file.  A missing or malformed file means "nothing decided"."""
    target = Path(path or DEFAULT_REVIEW_STATE_PATH)
    if not target.exists():
        return {}
    try:
        data = json.loads(target.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, UnicodeDecodeError):
        return {}
    if not isinstance(data, Mapping):
        return {}
    items = data.get("items", {})
    if not isinstance(items, Mapping):
        return {}
    return {
        str(key): dict(value)
        for key, value in items.items()
        if isinstance(value, Mapping) and value.get("status") in {STATUS_APPROVED, STATUS_IGNORED}
    }


def save_state(state: Mapping[str, Mapping[str, object]], path: Path | None = None) -> Path:
    target = Path(path or DEFAULT_REVIEW_STATE_PATH)
    target.parent.mkdir(parents=True, exist_ok=True)
    payload = {"version": STATE_VERSION, "items": {key: dict(value) for key, value in sorted(state.items())}}
    target.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return target


def _record_decision(
    item_id: str,
    status: str,
    action: str,
    details: Mapping[str, object] | None = None,
    path: Path | None = None,
    now: datetime | None = None,
) -> dict[str, dict[str, object]]:
    state = load_state(path)
    entry: dict[str, object] = {
        "status": status,
        "action": action,
        "at": (now or datetime.now()).replace(microsecond=0).isoformat(sep=" "),
    }
    entry.update({key: _jsonable(value) for key, value in (details or {}).items()})
    state[item_id] = entry
    save_state(state, path)
    return state


def _jsonable(value: object) -> object:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Mapping):
        return {str(key): _jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_jsonable(item) for item in value]
    return str(value)


def forget_decision(item_id: str, path: Path | None = None) -> dict[str, dict[str, object]]:
    """Put an ignored item back in the queue."""
    state = load_state(path)
    state.pop(item_id, None)
    save_state(state, path)
    return state


# --- suggestions -----------------------------------------------------------------


@dataclass(frozen=True)
class Suggestion:
    description: str
    amount: float | None
    transaction_type: str
    category: str
    transaction_date: date | None
    date_precision: str
    confidence: str
    rationale: str


def sheet_representative_dates(workbook_path: Path) -> dict[str, date]:
    """Midpoint of the dated rows on each sheet.

    Mirrors the stand-in date the migration already uses for a block with no dates,
    so an undated review row can be offered the same kind of date.  Read-only.
    """
    workbook = load_workbook(workbook_path, data_only=True)
    try:
        midpoints: dict[str, date] = {}
        for ws in workbook.worksheets:
            dates: list[date] = []
            for block in find_blocks(ws):
                column = block.date_column or block.description_column
                for row in range(block.header_row + 1, ws.max_row + 1):
                    value = ws.cell(row, column).value
                    if isinstance(value, datetime):
                        dates.append(value.date())
                    elif isinstance(value, date):
                        dates.append(value)
            if dates:
                ordinals = sorted(value.toordinal() for value in dates)
                midpoints[ws.title] = date.fromordinal(int(statistics.median_low(ordinals)))
        return midpoints
    finally:
        workbook.close()


def suggest(candidate: Candidate, sheet_date: date | None = None) -> Suggestion:
    """Prefill for one review row.  High confidence never means auto-approved."""
    note = candidate.review_note.lower()
    description = candidate.description
    transaction_type = candidate.transaction_type
    confidence = CONFIDENCE_LOW
    rationale = "Needs your judgement."

    is_formula = "formula" in note
    is_negative = "negative" in note

    if is_formula:
        rationale = "The amount cell is a formula, so it may combine several purchases or point at a total."
    elif is_negative and _RENT_CREDIT_NOTE in note:
        transaction_type = "Income"
        if "reimbursement" not in description.lower():
            description = f"{description} rent reimbursement"
        confidence = CONFIDENCE_HIGH
        rationale = "Negative amount in a Rent column: money coming back, recorded as positive Income."
    elif is_negative and _CREDIT_WORDS.search(candidate.description):
        transaction_type = "Income"
        confidence = CONFIDENCE_HIGH
        rationale = "Negative amount in an Expense column and the description names a credit or payout."
    elif is_negative:
        transaction_type = "Income"
        rationale = "Negative amount read as a credit, but the description does not say what it was."
    elif candidate.source_area == "Sidebar" and _INVESTMENT_WORDS.search(candidate.description):
        transaction_type = "Investment"
        confidence = CONFIDENCE_HIGH
        rationale = "Sidebar entry named after a brokerage."
    elif candidate.source_area == "Sidebar":
        rationale = "Value outside any transaction block: confirm what it is before importing."

    transaction_date = candidate.date
    precision = candidate.date_precision
    if transaction_date is None and sheet_date is not None:
        transaction_date = sheet_date
        precision = PRECISION_PERIOD

    category = categorize(description, transaction_type) if transaction_type == "Expense" else ""
    return Suggestion(
        description=description,
        amount=candidate.amount,
        transaction_type=transaction_type,
        category=category,
        transaction_date=transaction_date,
        date_precision=precision,
        confidence=confidence,
        rationale=rationale,
    )


# --- formulas --------------------------------------------------------------------


def formula_details(workbook_path: Path, sheet: str, description_reference: str) -> tuple[str, dict[str, str]]:
    """The amount cell's formula and the values of the cells it points at."""
    match = re.match(r"([A-Z]{1,3})(\d+)$", description_reference)
    if not match:
        return ("", {})
    amount_column = column_index_from_string(match.group(1)) + 1
    row = int(match.group(2))

    formulas = load_workbook(workbook_path, data_only=False)
    values = load_workbook(workbook_path, data_only=True)
    try:
        if sheet not in formulas.sheetnames:
            return ("", {})
        formula = formulas[sheet].cell(row, amount_column).value
        if not isinstance(formula, str) or not formula.startswith("="):
            return ("", {})

        referenced: dict[str, str] = {}
        for reference in dict.fromkeys(_CELL_REFERENCE.findall(formula)):
            plain = reference.replace("$", "")
            try:
                column = column_index_from_string(re.match(r"([A-Z]{1,3})", plain).group(1))
                target_row = int(re.search(r"(\d+)$", plain).group(1))
            except (AttributeError, ValueError):  # pragma: no cover - defensive
                continue
            value = values[sheet].cell(target_row, column).value
            own_formula = formulas[sheet].cell(target_row, column).value
            if isinstance(value, float) and not value.is_integer():
                text = f"{round(value, 2):.2f}"  # cached results carry float noise
            else:
                text = f"{value}"
            if isinstance(own_formula, str) and own_formula.startswith("="):
                text = f"{text} (itself {own_formula})"
            referenced[plain] = text
        return (formula, referenced)
    finally:
        formulas.close()
        values.close()


# --- queue -----------------------------------------------------------------------


@dataclass
class ReviewItem:
    item_id: str
    source_sheet: str
    source_reference: str
    source_period: str
    source_line: str
    source_area: str
    original_description: str
    original_amount: float | None
    original_type: str
    reason: str
    suggestion: Suggestion
    formula: str = ""
    formula_references: dict[str, str] = field(default_factory=dict)
    status: str = STATUS_UNRESOLVED
    decision: dict[str, object] = field(default_factory=dict)

    @property
    def amount_reference(self) -> str:
        match = re.match(r"([A-Z]{1,3})(\d+)$", self.source_reference)
        if not match:
            return self.source_reference
        return f"{get_column_letter(column_index_from_string(match.group(1)) + 1)}{match.group(2)}"

    @property
    def label(self) -> str:
        amount = f"${self.original_amount:,.2f}" if self.original_amount is not None else "no amount"
        return f"{self.original_description or '(no description)'} - {amount}"

    def as_row(self) -> dict[str, object]:
        return {
            "Item": self.item_id,
            "Source Period": self.source_period,
            "Source Sheet": self.source_sheet,
            "Source Reference": self.source_reference,
            "Source Line": self.source_line,
            "Original description": self.original_description,
            "Original amount": self.original_amount,
            "Reason": self.reason,
            "Suggested Description": self.suggestion.description,
            "Suggested Type": self.suggestion.transaction_type,
            "Suggested Category": self.suggestion.category,
            "Suggested Date": self.suggestion.transaction_date,
            "Suggested Date Precision": self.suggestion.date_precision,
            "Confidence": self.suggestion.confidence,
        }


@dataclass
class ReviewQueue:
    unresolved: list[ReviewItem]
    approved: list[ReviewItem]
    ignored: list[ReviewItem]
    already_in_tracker: list[str] = field(default_factory=list)

    @property
    def summary(self) -> dict[str, object]:
        return {
            "unresolved": len(self.unresolved),
            "unresolved_amount": round(
                sum(item.original_amount or 0.0 for item in self.unresolved), 2
            ),
            "approved": len(self.approved),
            "ignored": len(self.ignored),
            "high_confidence": sum(
                1 for item in self.unresolved if item.suggestion.confidence == CONFIDENCE_HIGH
            ),
        }

    def get(self, item_id: str) -> ReviewItem:
        for item in self.unresolved + self.approved + self.ignored:
            if item.item_id == item_id:
                return item
        raise KeyError(f"No review item {item_id!r}")


def _reason_for(candidate: Candidate) -> str:
    note = candidate.review_note
    parts = [part.strip() for part in note.split(";") if part.strip()]
    return parts[-1] if parts else "Held for review"


def item_id_for(candidate: Candidate) -> str:
    return f"{candidate.source_sheet}!{candidate.source_reference}"


def build_queue(
    source_workbook: Path,
    tracker_workbook: Path,
    state_path: Path | None = None,
    with_formulas: bool = True,
) -> ReviewQueue:
    """Current review queue: needs-review candidates minus anything already resolved.

    A row that is already in the tracker (matched on provenance) never comes back,
    so reloading the app after approving something is idempotent.
    """
    plan = dry_run(source_workbook, tracker_workbook)
    tracker = load_transactions(tracker_workbook)
    in_tracker = {
        identity
        for identity in (
            provenance_identity(record.get("Source Sheet"), record.get("Source Reference"))
            for record in tracker.to_dict("records")
        )
        if identity is not None
    }
    state = load_state(state_path)
    midpoints = sheet_representative_dates(source_workbook)

    unresolved: list[ReviewItem] = []
    approved: list[ReviewItem] = []
    ignored: list[ReviewItem] = []
    already: list[str] = []

    for candidate in plan.needs_review:
        item_id = item_id_for(candidate)
        suggestion = suggest(candidate, midpoints.get(candidate.source_sheet))
        item = ReviewItem(
            item_id=item_id,
            source_sheet=candidate.source_sheet,
            source_reference=candidate.source_reference,
            source_period=candidate.source_period,
            source_line=candidate.source_line,
            source_area=candidate.source_area,
            original_description=candidate.description,
            original_amount=candidate.amount,
            original_type=candidate.transaction_type,
            reason=_reason_for(candidate),
            suggestion=suggestion,
        )
        decision = state.get(item_id)
        provenance = (candidate.source_sheet, candidate.source_reference)
        if provenance in in_tracker:
            already.append(item_id)
            item.status = STATUS_APPROVED
            item.decision = dict(decision or {"status": STATUS_APPROVED, "action": "already in tracker"})
            approved.append(item)
            continue
        if decision and decision.get("status") == STATUS_APPROVED:
            item.status = STATUS_APPROVED
            item.decision = dict(decision)
            approved.append(item)
            continue
        if decision and decision.get("status") == STATUS_IGNORED:
            item.status = STATUS_IGNORED
            item.decision = dict(decision)
            ignored.append(item)
            continue

        if with_formulas and "formula" in candidate.review_note.lower():
            item.formula, item.formula_references = formula_details(
                source_workbook, candidate.source_sheet, candidate.source_reference
            )
        unresolved.append(item)

    unresolved.sort(key=lambda entry: -(entry.original_amount or 0.0))
    return ReviewQueue(
        unresolved=unresolved, approved=approved, ignored=ignored, already_in_tracker=already
    )


def queue_frame(items: list[ReviewItem]):
    import pandas as pd

    columns = [
        "Item",
        "Source Period",
        "Source Sheet",
        "Source Reference",
        "Original description",
        "Original amount",
        "Reason",
        "Suggested Description",
        "Suggested Type",
        "Suggested Category",
        "Suggested Date",
        "Suggested Date Precision",
        "Confidence",
    ]
    if not items:
        return pd.DataFrame(columns=columns)
    return pd.DataFrame([item.as_row() for item in items])[columns]


# --- actions ---------------------------------------------------------------------


class ReviewError(ValueError):
    """Raised when an approval would write something invalid or duplicated."""


def approve_item(
    item: ReviewItem,
    tracker_workbook: Path,
    overrides: Mapping[str, object] | None = None,
    state_path: Path | None = None,
    make_backup: bool = True,
) -> dict[str, object]:
    """Write one reviewed row to the tracker and mark the item resolved.

    ``overrides`` may set Description, Amount, Type, Category, Date and Date Precision;
    anything absent falls back to the suggestion.  Provenance is always preserved and
    never taken from the caller.
    """
    tracker_workbook = Path(tracker_workbook)
    overrides = dict(overrides or {})
    suggestion = item.suggestion

    description = str(overrides.get("Description", suggestion.description) or "").strip()
    transaction_type = str(overrides.get("Type", suggestion.transaction_type) or "").strip()
    raw_amount = overrides.get("Amount", suggestion.amount)
    category = str(overrides.get("Category", suggestion.category) or "").strip()
    transaction_date = overrides.get("Date", suggestion.transaction_date)
    precision = str(overrides.get("Date Precision", suggestion.date_precision) or "").strip()

    if not description:
        raise ReviewError("A description is required before importing this row.")
    if transaction_type not in VALID_TYPES:
        raise ReviewError(f"Type must be one of {', '.join(VALID_TYPES)}.")
    try:
        amount = round(float(raw_amount), 2)
    except (TypeError, ValueError):
        raise ReviewError("An amount is required before importing this row.") from None
    if amount <= 0:
        raise ReviewError("Amount must be greater than zero; amounts are stored positive.")
    if transaction_type != "Expense":
        category = ""
    if isinstance(transaction_date, datetime):
        transaction_date = transaction_date.date()
    if transaction_date is None:
        precision = PRECISION_UNKNOWN
    elif not precision or precision == PRECISION_UNKNOWN:
        precision = PRECISION_EXACT

    # Deterministic duplicate protection: one source cell, one transaction.
    existing = load_transactions(tracker_workbook)
    if not existing.empty:
        match = existing[
            (existing["Source Sheet"] == item.source_sheet)
            & (existing["Source Reference"] == item.source_reference)
        ]
        if not match.empty:
            _record_decision(
                item.item_id,
                STATUS_APPROVED,
                ACTION_APPROVE,
                {"duplicate": True, "note": "already present in the tracker"},
                state_path,
            )
            return {
                "imported": False,
                "duplicate": True,
                "item": item.item_id,
                "message": f"{item.item_id} is already in the tracker; nothing was written.",
            }

    row = {
        "Date": transaction_date,
        "Description": description,
        "Amount": amount,
        "Type": transaction_type,
        "Category": category,
        "Source Line": item.source_line,
        "Source Period": item.source_period,
        "Source Sheet": item.source_sheet,
        "Source Reference": item.source_reference,
        "Date Precision": precision,
    }

    backup = backup_workbook(tracker_workbook) if make_backup else None
    append_transactions(tracker_workbook, [row])

    edited = {
        key: value
        for key, value in {
            "Description": description if description != suggestion.description else None,
            "Amount": amount if suggestion.amount is None or amount != round(suggestion.amount, 2) else None,
            "Type": transaction_type if transaction_type != suggestion.transaction_type else None,
            # A category only counts as edited for Expense rows; for any other type the
            # code clears it itself, which is not a decision the user made.
            "Category": (
                category
                if transaction_type == "Expense" and category != suggestion.category
                else None
            ),
            "Date": transaction_date if transaction_date != suggestion.transaction_date else None,
            "Date Precision": precision if precision != suggestion.date_precision else None,
        }.items()
        if value is not None
    }
    action = ACTION_APPROVE_EDITED if edited else ACTION_APPROVE
    details: dict[str, object] = {"row": row, "confidence": suggestion.confidence}
    if edited:
        details["edited_fields"] = sorted(edited)
    if backup:
        details["backup"] = str(backup)
    _record_decision(item.item_id, STATUS_APPROVED, action, details, state_path)
    record_audit(
        tracker_workbook,
        AUDIT_APPROVE,
        {"item": item.item_id, "row": _jsonable(row), "edited_fields": sorted(edited)},
    )
    return {
        "imported": True,
        "duplicate": False,
        "item": item.item_id,
        "action": action,
        "row": row,
        "backup": str(backup) if backup else None,
        "message": f"Imported {description} (${amount:,.2f}) as {transaction_type}.",
    }


def ignore_item(
    item: ReviewItem | str,
    tracker_workbook: Path | None = None,
    note: str = "",
    state_path: Path | None = None,
) -> dict[str, object]:
    """Record that an item is intentionally excluded so it stops appearing."""
    item_id = item if isinstance(item, str) else item.item_id
    details: dict[str, object] = {"note": note} if note else {}
    if not isinstance(item, str):
        details.update(
            {
                "original_description": item.original_description,
                "original_amount": item.original_amount,
                "reason": item.reason,
            }
        )
    _record_decision(item_id, STATUS_IGNORED, ACTION_IGNORE, details, state_path)
    if tracker_workbook is not None:
        record_audit(Path(tracker_workbook), AUDIT_IGNORE, {"item": item_id, "note": note})
    return {"imported": False, "ignored": True, "item": item_id, "message": f"{item_id} will stay out."}


__all__ = [
    "ACTION_APPROVE",
    "ACTION_APPROVE_EDITED",
    "ACTION_IGNORE",
    "AUDIT_APPROVE",
    "AUDIT_IGNORE",
    "CONFIDENCE_HIGH",
    "CONFIDENCE_LOW",
    "DEFAULT_REVIEW_STATE_PATH",
    "STATUS_APPROVED",
    "STATUS_IGNORED",
    "STATUS_UNRESOLVED",
    "ReviewError",
    "ReviewItem",
    "ReviewQueue",
    "Suggestion",
    "approve_item",
    "build_queue",
    "forget_decision",
    "formula_details",
    "ignore_item",
    "item_id_for",
    "load_state",
    "queue_frame",
    "save_state",
    "sheet_representative_dates",
    "suggest",
]
