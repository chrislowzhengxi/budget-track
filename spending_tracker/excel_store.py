from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import shutil

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import pandas as pd

from spending_tracker.schema import (
    LEGACY_COLUMNS,
    PRECISION_EXACT,
    PRECISION_UNKNOWN,
    TRANSACTION_COLUMNS,
    empty_transactions_frame,
    normalize_transactions,
)


TRANSACTIONS_SHEET = "Transactions"
HEADERS = TRANSACTION_COLUMNS
BACKUP_DIRNAME = "backups"

_COLUMN_WIDTHS = {
    "Date": 12,
    "Description": 30,
    "Amount": 11,
    "Type": 12,
    "Category": 15,
    "Source Line": 32,
    "Source Period": 16,
    "Source Sheet": 20,
    "Source Reference": 16,
    "Date Precision": 14,
    "Added At": 20,
}


def load_transactions(workbook_path: Path) -> pd.DataFrame:
    """Read the Transactions sheet, tolerating pre-V3 (6 column) workbooks."""
    if not workbook_path.exists():
        return empty_transactions_frame()

    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if TRANSACTIONS_SHEET not in wb.sheetnames:
            return empty_transactions_frame()

        ws = wb[TRANSACTIONS_SHEET]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if not rows:
        return empty_transactions_frame()

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    records = [dict(zip(headers, row)) for row in rows[1:] if any(value is not None for value in row)]
    if not records:
        return empty_transactions_frame()

    return normalize_transactions(pd.DataFrame(records))


def ensure_workbook_copy(source_path: Path, output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if not output_path.exists():
        shutil.copy2(source_path, output_path)
    return output_path


def backup_workbook(workbook_path: Path, backup_dir: Path | None = None, now: datetime | None = None) -> Path | None:
    """Copy ``workbook_path`` to a timestamped file.  Returns None if nothing to back up."""
    if not workbook_path.exists():
        return None

    stamp = (now or datetime.now()).strftime("%Y%m%d-%H%M%S")
    target_dir = backup_dir or workbook_path.parent / BACKUP_DIRNAME
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"{workbook_path.stem}.backup-{stamp}{workbook_path.suffix}"
    counter = 2
    while target.exists():
        target = target_dir / f"{workbook_path.stem}.backup-{stamp}-{counter}{workbook_path.suffix}"
        counter += 1
    shutil.copy2(workbook_path, target)
    return target


def workbook_last_updated(workbook_path: Path) -> datetime | None:
    if not workbook_path.exists():
        return None
    return datetime.fromtimestamp(workbook_path.stat().st_mtime)


def upgrade_workbook_schema(workbook_path: Path) -> bool:
    """Widen an existing Transactions sheet to the current schema.

    Returns True when the workbook was changed.  Existing rows are preserved and
    remapped by header name, so pre-V3 tracker files keep every transaction.
    """
    if not workbook_path.exists():
        return False

    wb = load_workbook(workbook_path)
    try:
        if TRANSACTIONS_SHEET not in wb.sheetnames:
            return False
        ws = wb[TRANSACTIONS_SHEET]
        if not _needs_schema_upgrade(ws):
            return False
        _rewrite_with_schema(ws)
        wb.save(workbook_path)
    finally:
        wb.close()
    return True


def append_transactions(workbook_path: Path, rows: list[dict[str, object]]) -> int:
    if not rows:
        return 0

    wb = load_workbook(workbook_path)
    try:
        ws = _get_or_create_transactions_sheet(wb)
        added_at = datetime.now().replace(microsecond=0).isoformat(sep=" ")

        for row in rows:
            ws.append(_row_values(row, added_at))

        _format_sheet(ws)
        wb.save(workbook_path)
    finally:
        wb.close()
    return len(rows)


def replace_transactions(workbook_path: Path, frame: pd.DataFrame) -> int:
    """Overwrite the Transactions sheet with ``frame`` (used by edit/delete flows)."""
    normalized = normalize_transactions(frame)
    wb = load_workbook(workbook_path)
    try:
        ws = _get_or_create_transactions_sheet(wb)
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for record in normalized.to_dict("records"):
            ws.append(_row_values(record, record.get("Added At") or ""))
        _format_sheet(ws)
        wb.save(workbook_path)
    finally:
        wb.close()
    return len(normalized)


def _row_values(row: dict[str, object], added_at: str) -> list[object]:
    amount = row.get("Amount")
    values: list[object] = []
    for column in TRANSACTION_COLUMNS:
        if column == "Date":
            values.append(_normalize_excel_date(row.get("Date")))
        elif column == "Amount":
            values.append(float(amount) if amount is not None and amount == amount else None)
        elif column == "Added At":
            values.append(row.get("Added At") or added_at)
        elif column == "Date Precision":
            precision = row.get("Date Precision") or ""
            if not precision:
                precision = PRECISION_EXACT if row.get("Date") is not None else PRECISION_UNKNOWN
            values.append(precision)
        else:
            value = row.get(column, "")
            values.append("" if value is None or value != value else value)
    return values


def _normalize_excel_date(value: object) -> object:
    if value is None or value != value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if hasattr(value, "date") and callable(value.date):
        return value.date()
    return value


def _sheet_headers(ws) -> list[str]:
    if ws.max_row == 0:
        return []
    return [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]


def _needs_schema_upgrade(ws) -> bool:
    headers = _sheet_headers(ws)
    return headers[: len(TRANSACTION_COLUMNS)] != TRANSACTION_COLUMNS


def _rewrite_with_schema(ws) -> None:
    headers = _sheet_headers(ws)
    records = [
        dict(zip(headers, row))
        for row in ws.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in row)
    ]
    ws.delete_rows(1, max(ws.max_row, 1))
    ws.append(TRANSACTION_COLUMNS)
    for record in records:
        ws.append(_row_values(record, record.get("Added At") or ""))
    _format_sheet(ws)


def _get_or_create_transactions_sheet(wb):
    if TRANSACTIONS_SHEET in wb.sheetnames:
        ws = wb[TRANSACTIONS_SHEET]
        if ws.max_row == 0 or ws["A1"].value is None:
            ws.append(TRANSACTION_COLUMNS)
        elif _needs_schema_upgrade(ws):
            _rewrite_with_schema(ws)
        return ws

    ws = wb.create_sheet(TRANSACTIONS_SHEET)
    ws.append(TRANSACTION_COLUMNS)
    _format_sheet(ws)
    return ws


def _format_sheet(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for index, column in enumerate(TRANSACTION_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = _COLUMN_WIDTHS.get(column, 14)

    amount_col = TRANSACTION_COLUMNS.index("Amount") + 1
    for row in ws.iter_rows(min_row=2, min_col=amount_col, max_col=amount_col):
        for cell in row:
            cell.number_format = "0.00"

    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.number_format = "yyyy-mm-dd"

    last_column = get_column_letter(len(TRANSACTION_COLUMNS))
    ref = f"A1:{last_column}{max(ws.max_row, 2)}"
    if "TransactionsTable" not in ws.tables:
        table = Table(displayName="TransactionsTable", ref=ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
    else:
        ws.tables["TransactionsTable"].ref = ref


__all__ = [
    "BACKUP_DIRNAME",
    "HEADERS",
    "LEGACY_COLUMNS",
    "TRANSACTIONS_SHEET",
    "append_transactions",
    "backup_workbook",
    "empty_transactions_frame",
    "ensure_workbook_copy",
    "load_transactions",
    "replace_transactions",
    "upgrade_workbook_schema",
    "workbook_last_updated",
]
