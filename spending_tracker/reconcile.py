"""Reconcile the original historical workbook against the tracker's Transactions sheet.

This is an audit tool, not part of the app.  It answers one question: for every
Source Period and Type, does the tracker hold what the original workbook holds, and
if not, exactly which cells are missing and why.

Three independent layers are compared so a bug in the extractor cannot hide:

1. raw amount-column cell sums read straight from the workbook (what Excel's own
   ``SUM(D:D)`` sees, including cached formula results);
2. the candidate rows the extractor produces, with their status;
3. the rows actually present in the tracker, matched by provenance (sheet + cell).
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
import csv
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
import pandas as pd

from spending_tracker.excel_store import load_transactions
from spending_tracker.historical import (
    STATUS_NEEDS_REVIEW,
    STATUS_READY,
    STATUS_SKIPPED,
    Candidate,
    extract,
    find_blocks,
)
from spending_tracker.migration import candidate_provenance, provenance_identity


RECONCILIATION_CSV = "historical_reconciliation.csv"
RECONCILIATION_MD = "historical_reconciliation.md"

RECONCILIATION_COLUMNS = [
    "Source Period",
    "Source Sheet",
    "Block",
    "Type",
    "Original rows",
    "Original amount",
    "Migrated rows",
    "Migrated amount",
    "Excluded rows",
    "Excluded amount",
    "Difference",
    "Exclusion reasons",
]

# Buckets for rows held back from migration.  Block rows are classified by the
# structural reason the extractor recorded; sidebar rows are only ever hints.
BUCKET_FORMULA = "Formula amount"
BUCKET_RENT_CREDIT = "Likely Rent reimbursement (negative in a Rent column)"
BUCKET_INCOME_CREDIT = "Likely real Income (negative in an Expense column)"
BUCKET_NO_DATE = "Likely real transaction, no date in its block"
BUCKET_AMBIGUOUS = "Genuinely ambiguous"
BUCKET_SIDEBAR_PLAN = "Sidebar: inside a summed range (budget or plan figure)"
BUCKET_SIDEBAR_INCOME = "Sidebar: likely real Income"
BUCKET_SIDEBAR_INVESTMENT = "Sidebar: likely Investment"
BUCKET_SIDEBAR_RENT = "Sidebar: likely Rent or utilities"
BUCKET_SIDEBAR_REFERENCE = "Sidebar: reference figure, not a transaction"
BUCKET_SIDEBAR_OTHER = "Sidebar: likely Expense or unclear"

BUCKET_ORDER = [
    BUCKET_NO_DATE,
    BUCKET_FORMULA,
    BUCKET_INCOME_CREDIT,
    BUCKET_RENT_CREDIT,
    BUCKET_SIDEBAR_INCOME,
    BUCKET_SIDEBAR_INVESTMENT,
    BUCKET_SIDEBAR_RENT,
    BUCKET_SIDEBAR_OTHER,
    BUCKET_SIDEBAR_PLAN,
    BUCKET_SIDEBAR_REFERENCE,
    BUCKET_AMBIGUOUS,
]

_INCOME_WORDS = re.compile(r"deposit|refund|reimburs|credit|cashback|cash back|payout|stipend|rebate", re.I)
_INVESTMENT_WORDS = re.compile(r"invest|ibkr|robinhood|brokerage|vanguard|fidelity|schwab", re.I)
_RENT_WORDS = re.compile(r"\brent\b|utils|utilities", re.I)
_REFERENCE_WORDS = re.compile(r"price|estimate|actual|plan|opening|target amount", re.I)


@dataclass
class BlockTotal:
    """Raw amount-column arithmetic for one block, independent of the extractor."""

    sheet: str
    block: str
    transaction_type: str
    period: str
    cells: int = 0
    total: float = 0.0
    literal_cells: int = 0
    literal_total: float = 0.0
    formula_cells: int = 0
    formula_total: float = 0.0
    negative_cells: int = 0
    negative_total: float = 0.0

    @property
    def positive_literal_total(self) -> float:
        return round(self.literal_total - self.negative_total, 2)


def summed_cells(ws) -> set[tuple[int, int]]:
    """Cells covered by a SUM range in some formula on the sheet.

    A sidebar number that a formula adds up is part of a total or a budget table,
    not a transaction, so it is bucketed separately.  Only bounded ranges count:
    whole-column references such as ``SUM(D:D)`` always point at a transaction block
    column, which the sidebar scan never inspects anyway.
    """
    covered: set[tuple[int, int]] = set()
    for row in ws.iter_rows():
        for cell in row:
            value = cell.value
            if not isinstance(value, str) or not value.startswith("="):
                continue
            for reference in re.findall(r"(?<![A-Z0-9_])(\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+)", value):
                try:
                    min_col, min_row, max_col, max_row = range_boundaries(reference.replace("$", ""))
                except ValueError:  # pragma: no cover - defensive
                    continue
                if None in (min_col, min_row, max_col, max_row):
                    continue
                if (max_row - min_row) > 2000 or (max_col - min_col) > 50:
                    continue
                for r in range(min_row, max_row + 1):
                    for c in range(min_col, max_col + 1):
                        if (r, c) != (cell.row, cell.column):
                            covered.add((r, c))
    return covered


def block_totals(workbook_path: Path) -> list[BlockTotal]:
    """Sum every numeric cell in each block's amount column, straight from the file."""
    values = load_workbook(workbook_path, data_only=True)
    formulas = load_workbook(workbook_path, data_only=False)
    try:
        totals: list[BlockTotal] = []
        for ws in values.worksheets:
            ws_formulas = formulas[ws.title]
            for block in find_blocks(ws):
                total = BlockTotal(
                    sheet=ws.title,
                    block=(
                        f"{get_column_letter(block.description_column)}/"
                        f"{get_column_letter(block.amount_column)}"
                    ),
                    transaction_type=block.transaction_type,
                    period=block.period,
                )
                for row in range(block.header_row + 1, ws.max_row + 1):
                    value = ws.cell(row, block.amount_column).value
                    if not isinstance(value, (int, float)) or isinstance(value, bool):
                        continue
                    formula = ws_formulas.cell(row, block.amount_column).value
                    total.cells += 1
                    total.total = round(total.total + value, 2)
                    if isinstance(formula, str) and formula.startswith("="):
                        total.formula_cells += 1
                        total.formula_total = round(total.formula_total + value, 2)
                    else:
                        total.literal_cells += 1
                        total.literal_total = round(total.literal_total + value, 2)
                    if value < 0:
                        total.negative_cells += 1
                        total.negative_total = round(total.negative_total + value, 2)
                totals.append(total)
        return totals
    finally:
        values.close()
        formulas.close()


def review_bucket(candidate: Candidate, summed: set[tuple[int, int]] | None = None) -> str:
    """Which bucket a held-back row belongs in."""
    note = candidate.review_note.lower()
    if candidate.source_area == "Block":
        if "formula" in note:
            return BUCKET_FORMULA
        if "negative" in note:
            return BUCKET_RENT_CREDIT if "'rent' column" in note else BUCKET_INCOME_CREDIT
        if "no date information anywhere" in note:
            return BUCKET_NO_DATE
        return BUCKET_AMBIGUOUS

    if summed and _cell_of(candidate) in summed:
        return BUCKET_SIDEBAR_PLAN
    description = candidate.description
    if _REFERENCE_WORDS.search(description):
        return BUCKET_SIDEBAR_REFERENCE
    if _INVESTMENT_WORDS.search(description):
        return BUCKET_SIDEBAR_INVESTMENT
    if _INCOME_WORDS.search(description):
        return BUCKET_SIDEBAR_INCOME
    if _RENT_WORDS.search(description):
        return BUCKET_SIDEBAR_RENT
    return BUCKET_SIDEBAR_OTHER


def _cell_of(candidate: Candidate) -> tuple[int, int]:
    match = re.match(r"([A-Z]{1,3})(\d+)$", candidate.source_reference)
    if not match:
        return (0, 0)
    column = 0
    for character in match.group(1):
        column = column * 26 + (ord(character) - ord("A") + 1)
    return (int(match.group(2)), column)


def _reason(candidate: Candidate) -> str:
    note = candidate.review_note.lower()
    if "formula" in note:
        return "formula amount"
    if "negative" in note:
        return "negative amount read as a credit"
    if "no date information anywhere" in note:
        return "no date available in its block"
    if "total, heading or label" in note:
        return "looks like a total or label"
    if "no amount" in note:
        return "no amount"
    if "zero" in note:
        return "zero amount"
    if "no description" in note:
        return "no description"
    if candidate.source_area == "Sidebar":
        return "outside any transaction block"
    return "other"


@dataclass
class Reconciliation:
    rows: list[dict[str, object]]
    totals_by_type: dict[str, dict[str, float]]
    block_totals: list[BlockTotal]
    buckets: dict[str, list[Candidate]] = field(default_factory=dict)
    candidates: list[Candidate] = field(default_factory=list)
    tracker_rows: int = 0
    tracker_migrated_rows: int = 0
    tracker_manual_rows: int = 0
    source_workbook: str = ""
    tracker_workbook: str = ""


def reconcile(source_workbook: Path, tracker_workbook: Path) -> Reconciliation:
    """Compare the original workbook with the tracker, period by period and type by type."""
    extraction = extract(source_workbook)
    tracker = load_transactions(tracker_workbook)
    migrated_keys = {
        identity
        for identity in (
            provenance_identity(record.get("Source Sheet"), record.get("Source Reference"))
            for record in tracker.to_dict("records")
        )
        if identity is not None
    }

    grouped: dict[tuple[str, str, str, str], dict[str, object]] = {}
    block_lookup = {
        (block.sheet, block.period, block.transaction_type): (
            f"{get_column_letter(block.description_column)}/{get_column_letter(block.amount_column)}"
        )
        for block in extraction.blocks
    }

    for candidate in extraction.candidates:
        if candidate.source_area != "Block":
            continue
        block_type = _block_type_of(candidate, extraction.blocks)
        key = (candidate.source_period, candidate.source_sheet, block_type, candidate.transaction_type)
        entry = grouped.setdefault(
            key,
            {
                "Source Period": candidate.source_period,
                "Source Sheet": candidate.source_sheet,
                "Block": block_lookup.get(
                    (candidate.source_sheet, candidate.source_period, block_type), ""
                ),
                "Type": candidate.transaction_type,
                "Original rows": 0,
                "Original amount": 0.0,
                "Migrated rows": 0,
                "Migrated amount": 0.0,
                "Excluded rows": 0,
                "Excluded amount": 0.0,
                "Difference": 0.0,
                "_reasons": defaultdict(lambda: [0, 0.0]),
            },
        )
        amount = candidate.amount or 0.0
        entry["Original rows"] += 1
        entry["Original amount"] = round(entry["Original amount"] + amount, 2)
        if candidate_provenance(candidate) in migrated_keys:
            entry["Migrated rows"] += 1
            entry["Migrated amount"] = round(entry["Migrated amount"] + amount, 2)
        else:
            entry["Excluded rows"] += 1
            entry["Excluded amount"] = round(entry["Excluded amount"] + amount, 2)
            reason = entry["_reasons"][_reason(candidate)]
            reason[0] += 1
            reason[1] = round(reason[1] + amount, 2)

    rows: list[dict[str, object]] = []
    totals_by_type: dict[str, dict[str, float]] = {}
    for key in sorted(grouped, key=lambda item: (item[3], item[0], item[1])):
        entry = grouped[key]
        entry["Difference"] = round(entry["Original amount"] - entry["Migrated amount"], 2)
        entry["Exclusion reasons"] = "; ".join(
            f"{name} ({count} row(s), ${amount:,.2f})"
            for name, (count, amount) in sorted(entry.pop("_reasons").items())
        )
        rows.append({column: entry[column] for column in RECONCILIATION_COLUMNS})
        bucket = totals_by_type.setdefault(
            str(entry["Type"]),
            {
                "original_rows": 0,
                "original_amount": 0.0,
                "migrated_rows": 0,
                "migrated_amount": 0.0,
                "excluded_rows": 0,
                "excluded_amount": 0.0,
            },
        )
        bucket["original_rows"] += int(entry["Original rows"])
        bucket["original_amount"] = round(bucket["original_amount"] + float(entry["Original amount"]), 2)
        bucket["migrated_rows"] += int(entry["Migrated rows"])
        bucket["migrated_amount"] = round(bucket["migrated_amount"] + float(entry["Migrated amount"]), 2)
        bucket["excluded_rows"] += int(entry["Excluded rows"])
        bucket["excluded_amount"] = round(bucket["excluded_amount"] + float(entry["Excluded amount"]), 2)

    summed: set[tuple[int, int]] = set()
    values = load_workbook(source_workbook, data_only=False)
    try:
        per_sheet = {ws.title: summed_cells(ws) for ws in values.worksheets}
    finally:
        values.close()

    buckets: dict[str, list[Candidate]] = defaultdict(list)
    for candidate in extraction.candidates:
        if candidate.status == STATUS_READY and candidate_provenance(candidate) in migrated_keys:
            continue
        if candidate.status == STATUS_SKIPPED:
            continue
        buckets[review_bucket(candidate, per_sheet.get(candidate.source_sheet, summed))].append(candidate)

    return Reconciliation(
        rows=rows,
        totals_by_type=totals_by_type,
        block_totals=block_totals(source_workbook),
        buckets=dict(buckets),
        candidates=extraction.candidates,
        tracker_rows=len(tracker),
        tracker_migrated_rows=int((tracker["Source Sheet"] != "").sum()) if not tracker.empty else 0,
        tracker_manual_rows=int((tracker["Source Sheet"] == "").sum()) if not tracker.empty else 0,
        source_workbook=str(source_workbook),
        tracker_workbook=str(tracker_workbook),
    )


def _block_type_of(candidate: Candidate, blocks) -> str:
    """The type of the block a candidate came from (before any re-typing)."""
    for block in blocks:
        if block.sheet == candidate.source_sheet and block.period == candidate.source_period:
            reference = _cell_of(candidate)
            if reference[1] == block.description_column:
                return block.transaction_type
    return candidate.transaction_type


def write_reconciliation_csv(result: Reconciliation, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=RECONCILIATION_COLUMNS)
        writer.writeheader()
        writer.writerows(result.rows)
    return path


def write_reconciliation_markdown(result: Reconciliation, path: Path) -> Path:
    lines: list[str] = []
    lines.append("# Historical reconciliation report")
    lines.append("")
    lines.append(f"- Generated: {datetime.now().replace(microsecond=0).isoformat(sep=' ')}")
    lines.append(f"- Original workbook (read-only): `{result.source_workbook}`")
    lines.append(f"- Tracker workbook: `{result.tracker_workbook}`")
    lines.append(
        f"- Tracker holds {result.tracker_rows} row(s): "
        f"{result.tracker_migrated_rows} migrated, {result.tracker_manual_rows} entered by hand"
    )
    lines.append("")

    lines.append("## 1. Raw workbook arithmetic per block")
    lines.append("")
    lines.append("Amount-column cells read straight from the file, before any interpretation.")
    lines.append("")
    lines.append("| Sheet | Block | Type | Period | Cells | Sum | Literal | Formula | Negative |")
    lines.append("| --- | --- | --- | --- | ---: | ---: | ---: | ---: | ---: |")
    for total in result.block_totals:
        lines.append(
            f"| {total.sheet} | {total.block} | {total.transaction_type} | {total.period} | "
            f"{total.cells} | ${total.total:,.2f} | {total.literal_cells} / ${total.literal_total:,.2f} | "
            f"{total.formula_cells} / ${total.formula_total:,.2f} | "
            f"{total.negative_cells} / ${total.negative_total:,.2f} |"
        )
    lines.append("")

    lines.append("## 2. Original vs migrated, per Source Period and Type")
    lines.append("")
    lines.append(
        "| Source Period | Type | Orig rows | Orig $ | Migr rows | Migr $ | Excl rows | Excl $ | Diff $ | Reasons |"
    )
    lines.append("| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | --- |")
    for row in result.rows:
        lines.append(
            f"| {row['Source Period']} | {row['Type']} | {row['Original rows']} | "
            f"${row['Original amount']:,.2f} | {row['Migrated rows']} | ${row['Migrated amount']:,.2f} | "
            f"{row['Excluded rows']} | ${row['Excluded amount']:,.2f} | ${row['Difference']:,.2f} | "
            f"{row['Exclusion reasons'] or '-'} |"
        )
    lines.append("")

    lines.append("## 3. Totals by Type")
    lines.append("")
    lines.append("| Type | Orig rows | Orig $ | Migr rows | Migr $ | Excl rows | Excl $ |")
    lines.append("| --- | ---: | ---: | ---: | ---: | ---: | ---: |")
    for name, bucket in sorted(result.totals_by_type.items()):
        lines.append(
            f"| {name} | {bucket['original_rows']} | ${bucket['original_amount']:,.2f} | "
            f"{bucket['migrated_rows']} | ${bucket['migrated_amount']:,.2f} | "
            f"{bucket['excluded_rows']} | ${bucket['excluded_amount']:,.2f} |"
        )
    lines.append("")

    lines.append("## 4. Rows held back from migration, by bucket")
    lines.append("")
    held = sum(len(items) for items in result.buckets.values())
    held_amount = sum(
        candidate.amount or 0.0 for items in result.buckets.values() for candidate in items
    )
    lines.append(f"{held} row(s), ${held_amount:,.2f} in total.")
    lines.append("")
    for name in BUCKET_ORDER:
        items = result.buckets.get(name, [])
        if not items:
            continue
        total = sum(candidate.amount or 0.0 for candidate in items)
        lines.append(f"### {name} - {len(items)} row(s), ${total:,.2f}")
        lines.append("")
        lines.append("| Source | Description | Amount | Type | Period | Date precision | Note |")
        lines.append("| --- | --- | ---: | --- | --- | --- | --- |")
        for candidate in sorted(items, key=lambda item: -(item.amount or 0.0)):
            amount = f"${candidate.amount:,.2f}" if candidate.amount is not None else "(none)"
            lines.append(
                f"| `{candidate.source_sheet}!{candidate.source_reference}` | {candidate.description} | "
                f"{amount} | {candidate.transaction_type} | {candidate.source_period} | "
                f"{candidate.date_precision} | {candidate.review_note} |"
            )
        lines.append("")

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


__all__ = [
    "BUCKET_AMBIGUOUS",
    "BUCKET_FORMULA",
    "BUCKET_INCOME_CREDIT",
    "BUCKET_NO_DATE",
    "BUCKET_ORDER",
    "BUCKET_RENT_CREDIT",
    "BUCKET_SIDEBAR_INCOME",
    "BUCKET_SIDEBAR_INVESTMENT",
    "BUCKET_SIDEBAR_OTHER",
    "BUCKET_SIDEBAR_PLAN",
    "BUCKET_SIDEBAR_REFERENCE",
    "BUCKET_SIDEBAR_RENT",
    "BlockTotal",
    "RECONCILIATION_COLUMNS",
    "RECONCILIATION_CSV",
    "RECONCILIATION_MD",
    "Reconciliation",
    "block_totals",
    "reconcile",
    "review_bucket",
    "summed_cells",
    "write_reconciliation_csv",
    "write_reconciliation_markdown",
]
