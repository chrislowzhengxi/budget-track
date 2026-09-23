from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pandas as pd

from spending_tracker.excel_store import (
    HEADERS,
    TRANSACTIONS_SHEET,
    append_transactions,
    backup_workbook,
    empty_transactions_frame,
    ensure_workbook_copy,
    load_transactions,
    replace_transactions,
    upgrade_workbook_schema,
    workbook_last_updated,
)
from spending_tracker.schema import PRECISION_EXACT, TRANSACTION_COLUMNS


def _create_source_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Existing"
    sheet["A1"] = "Do not touch"
    sheet["B2"] = 42
    workbook.save(path)


def test_ensure_workbook_copy_creates_parent_and_preserves_source(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "nested" / "copy.xlsx"
    _create_source_workbook(source)

    result = ensure_workbook_copy(source, output)

    assert result == output
    assert output.exists()
    source_wb = load_workbook(source)
    output_wb = load_workbook(output)
    assert source_wb.sheetnames == ["Existing"]
    assert output_wb.sheetnames == ["Existing"]
    assert output_wb["Existing"]["A1"].value == "Do not touch"


def test_ensure_workbook_copy_does_not_overwrite_existing_output(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "copy.xlsx"
    _create_source_workbook(source)
    _create_source_workbook(output)
    output_wb = load_workbook(output)
    output_wb["Existing"]["A1"] = "Already edited"
    output_wb.save(output)

    ensure_workbook_copy(source, output)

    reloaded = load_workbook(output)
    assert reloaded["Existing"]["A1"].value == "Already edited"


def test_append_transactions_creates_transactions_sheet_without_changing_existing_sheet(tmp_path: Path) -> None:
    workbook_path = tmp_path / "spending.xlsx"
    _create_source_workbook(workbook_path)

    rows = [
        {
            "Date": date(2026, 9, 14),
            "Description": "Uber",
            "Amount": 12.0,
            "Type": "Expense",
            "Source Line": "uber 12",
        },
        {
            "Date": date(2026, 9, 14),
            "Description": "TA",
            "Amount": 695.56,
            "Type": "Income",
            "Source Line": "TA (695.56)",
        },
    ]

    count = append_transactions(workbook_path, rows)

    assert count == 2
    workbook = load_workbook(workbook_path)
    assert workbook["Existing"]["A1"].value == "Do not touch"
    assert TRANSACTIONS_SHEET in workbook.sheetnames

    sheet = workbook[TRANSACTIONS_SHEET]
    assert [sheet.cell(1, column).value for column in range(1, len(HEADERS) + 1)] == HEADERS
    assert sheet["A2"].value == datetime(2026, 9, 14)
    assert [sheet.cell(2, column).value for column in (2, 3, 4)] == ["Uber", 12, "Expense"]
    assert sheet.cell(2, HEADERS.index("Source Line") + 1).value == "uber 12"
    assert sheet["A3"].value == datetime(2026, 9, 14)
    assert [sheet.cell(3, column).value for column in (2, 3, 4)] == ["TA", 695.56, "Income"]
    assert sheet.cell(3, HEADERS.index("Source Line") + 1).value == "TA (695.56)"
    assert sheet.cell(2, HEADERS.index("Added At") + 1).value
    assert sheet["A2"].number_format == "yyyy-mm-dd"
    assert sheet["C2"].number_format == "0.00"
    assert "TransactionsTable" in sheet.tables
    assert sheet.tables["TransactionsTable"].ref == "A1:K3"


def test_append_transactions_appends_to_existing_transactions_sheet_and_expands_table(tmp_path: Path) -> None:
    workbook_path = tmp_path / "spending.xlsx"
    _create_source_workbook(workbook_path)

    first_count = append_transactions(
        workbook_path,
        [
            {
                "Date": date(2026, 9, 14),
                "Description": "Rent",
                "Amount": 1800,
                "Type": "Rent",
                "Source Line": "rent 1800",
            }
        ],
    )
    second_count = append_transactions(
        workbook_path,
        [
            {
                "Date": date(2026, 9, 15),
                "Description": "Investment",
                "Amount": 500,
                "Type": "Investment",
                "Source Line": "investment 500",
            }
        ],
    )

    workbook = load_workbook(workbook_path)
    sheet = workbook[TRANSACTIONS_SHEET]
    assert first_count == 1
    assert second_count == 1
    assert sheet.max_row == 3
    assert sheet["B2"].value == "Rent"
    assert sheet["B3"].value == "Investment"
    assert sheet.tables["TransactionsTable"].ref == "A1:K3"


def test_append_transactions_with_no_rows_does_not_create_sheet(tmp_path: Path) -> None:
    workbook_path = tmp_path / "spending.xlsx"
    _create_source_workbook(workbook_path)

    count = append_transactions(workbook_path, [])

    workbook = load_workbook(workbook_path)
    assert count == 0
    assert TRANSACTIONS_SHEET not in workbook.sheetnames


def test_load_transactions_returns_empty_frame_when_workbook_is_missing(tmp_path: Path) -> None:
    transactions = load_transactions(tmp_path / "missing.xlsx")

    assert list(transactions.columns) == HEADERS
    assert transactions.empty


def test_load_transactions_returns_empty_frame_when_sheet_is_missing(tmp_path: Path) -> None:
    workbook_path = tmp_path / "spending.xlsx"
    _create_source_workbook(workbook_path)

    transactions = load_transactions(workbook_path)

    assert list(transactions.columns) == HEADERS
    assert transactions.empty


def test_load_transactions_reads_dates_and_numeric_amounts_without_modifying_workbook(tmp_path: Path) -> None:
    workbook_path = tmp_path / "spending.xlsx"
    _create_source_workbook(workbook_path)
    before_mtime = workbook_path.stat().st_mtime_ns
    append_transactions(
        workbook_path,
        [
            {
                "Date": date(2026, 9, 14),
                "Description": "Uber",
                "Amount": 12,
                "Type": "Expense",
                "Source Line": "uber 12",
            }
        ],
    )
    after_write_mtime = workbook_path.stat().st_mtime_ns

    transactions = load_transactions(workbook_path)
    after_read_mtime = workbook_path.stat().st_mtime_ns

    assert after_write_mtime >= before_mtime
    assert after_read_mtime == after_write_mtime
    assert len(transactions) == 1
    assert transactions.loc[0, "Date"] == pd.Timestamp("2026-09-14")
    assert transactions.loc[0, "Amount"] == 12
    assert transactions.loc[0, "Description"] == "Uber"
    assert transactions.loc[0, "Type"] == "Expense"


def test_load_transactions_reads_a_pre_v3_six_column_workbook(legacy_tracker: Path) -> None:
    transactions = load_transactions(legacy_tracker)

    assert list(transactions.columns) == TRANSACTION_COLUMNS
    assert len(transactions) == 5
    assert transactions.loc[0, "Description"] == "Uber"
    assert transactions.loc[0, "Amount"] == 12
    assert transactions.loc[0, "Source Line"] == "uber 12"
    # Hand-entered rows carried a real date, so their precision is Exact and the
    # new provenance columns are simply blank.
    assert set(transactions["Date Precision"]) == {PRECISION_EXACT}
    assert set(transactions["Category"]) == {""}
    assert set(transactions["Source Sheet"]) == {""}


def test_upgrade_workbook_schema_widens_the_sheet_and_keeps_every_row(legacy_tracker: Path) -> None:
    before = load_transactions(legacy_tracker)

    changed = upgrade_workbook_schema(legacy_tracker)

    assert changed is True
    workbook = load_workbook(legacy_tracker)
    sheet = workbook[TRANSACTIONS_SHEET]
    assert [cell.value for cell in sheet[1]] == TRANSACTION_COLUMNS
    assert sheet.max_row == 6
    assert workbook["24-25 Fall, Winter"]["B1"].value == "Fall 24"
    after = load_transactions(legacy_tracker)
    pd.testing.assert_frame_equal(before, after)
    assert upgrade_workbook_schema(legacy_tracker) is False


def test_appending_to_a_legacy_workbook_upgrades_it_without_losing_rows(legacy_tracker: Path) -> None:
    count = append_transactions(
        legacy_tracker,
        [
            {
                "Date": date(2026, 9, 20),
                "Description": "Tatte",
                "Amount": 7.85,
                "Type": "Expense",
                "Category": "Dining",
                "Source Line": "tatte 7.85",
            }
        ],
    )

    transactions = load_transactions(legacy_tracker)
    assert count == 1
    assert len(transactions) == 6
    assert transactions.iloc[-1]["Category"] == "Dining"
    assert transactions.iloc[0]["Description"] == "Uber"
    workbook = load_workbook(legacy_tracker)
    assert [cell.value for cell in workbook[TRANSACTIONS_SHEET][1]] == TRANSACTION_COLUMNS


def test_append_transactions_accepts_rows_without_the_new_columns(tmp_path: Path) -> None:
    workbook_path = tmp_path / "spending.xlsx"
    _create_source_workbook(workbook_path)

    append_transactions(
        workbook_path,
        [{"Date": date(2026, 9, 14), "Description": "Uber", "Amount": 12.0, "Type": "Expense"}],
    )

    transactions = load_transactions(workbook_path)
    assert transactions.loc[0, "Source Line"] == ""
    assert transactions.loc[0, "Date Precision"] == PRECISION_EXACT
    assert transactions.loc[0, "Added At"]


def test_replace_transactions_overwrites_rows_and_keeps_other_sheets(legacy_tracker: Path) -> None:
    transactions = load_transactions(legacy_tracker)
    kept = transactions[transactions["Description"] != "Uber"].copy()

    written = replace_transactions(legacy_tracker, kept)

    reloaded = load_transactions(legacy_tracker)
    assert written == 4
    assert len(reloaded) == 4
    assert "Uber" not in set(reloaded["Description"])
    workbook = load_workbook(legacy_tracker)
    assert workbook["24-25 Fall, Winter"]["B1"].value == "Fall 24"
    assert workbook[TRANSACTIONS_SHEET].tables["TransactionsTable"].ref == "A1:K5"


def test_replace_transactions_can_empty_the_sheet(legacy_tracker: Path) -> None:
    written = replace_transactions(legacy_tracker, empty_transactions_frame())

    assert written == 0
    assert load_transactions(legacy_tracker).empty
    workbook = load_workbook(legacy_tracker)
    assert [cell.value for cell in workbook[TRANSACTIONS_SHEET][1]] == TRANSACTION_COLUMNS


def test_backup_workbook_writes_a_timestamped_copy(legacy_tracker: Path) -> None:
    backup = backup_workbook(legacy_tracker, now=datetime(2026, 9, 23, 10, 30, 0))

    assert backup is not None
    assert backup.name == "tracker-legacy.backup-20260923-103000.xlsx"
    assert backup.read_bytes() == legacy_tracker.read_bytes()

    second = backup_workbook(legacy_tracker, now=datetime(2026, 9, 23, 10, 30, 0))
    assert second is not None
    assert second != backup


def test_backup_workbook_returns_none_for_a_missing_file(tmp_path: Path) -> None:
    assert backup_workbook(tmp_path / "missing.xlsx") is None


def test_workbook_last_updated_reports_the_file_mtime(legacy_tracker: Path) -> None:
    updated = workbook_last_updated(legacy_tracker)

    assert updated is not None
    assert abs(updated.timestamp() - legacy_tracker.stat().st_mtime) < 1
    assert workbook_last_updated(legacy_tracker.parent / "missing.xlsx") is None
