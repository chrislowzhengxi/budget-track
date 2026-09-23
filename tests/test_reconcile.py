from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook

from spending_tracker.excel_store import load_transactions
from spending_tracker.migration import run_migration
from spending_tracker.reconcile import (
    BUCKET_FORMULA,
    BUCKET_INCOME_CREDIT,
    BUCKET_NO_DATE,
    BUCKET_RENT_CREDIT,
    BUCKET_SIDEBAR_INCOME,
    BUCKET_SIDEBAR_PLAN,
    RECONCILIATION_COLUMNS,
    block_totals,
    reconcile,
    summed_cells,
    write_reconciliation_csv,
    write_reconciliation_markdown,
)


def test_block_totals_match_the_raw_workbook_arithmetic(historical_workbook: Path) -> None:
    totals = {(total.sheet, total.block): total for total in block_totals(historical_workbook)}

    fall = totals[("23-24 Fall, Winter", "C/D")]
    # Every numeric cell in the amount column, exactly as SUM(D:D) sees it:
    # 12 + 31.5 + 20 + 12 - 25 (credit) + 100 (a total row) + 99 (no description) + 0.
    # The =10+5 formula has no cached value in a file openpyxl wrote, so it adds nothing.
    assert fall.cells == 8
    assert fall.literal_cells == 8
    assert fall.formula_cells == 0
    assert fall.negative_cells == 1
    assert fall.negative_total == -25
    assert fall.total == 249.5
    assert fall.positive_literal_total == 274.5

    rent = totals[("Rent, Income", "B/C")]
    assert rent.cells == 3
    assert rent.total == 775.0  # 1400 + 75 - 700
    assert rent.negative_total == -700
    assert rent.positive_literal_total == 1475.0


def test_reconciliation_shows_zero_difference_after_a_full_migration(
    historical_workbook: Path, legacy_tracker: Path
) -> None:
    run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)

    result = reconcile(historical_workbook, legacy_tracker)

    by_type = result.totals_by_type
    assert by_type["Expense"]["migrated_rows"] == 6
    assert by_type["Rent"]["migrated_rows"] == 2
    assert by_type["Income"]["migrated_rows"] == 1
    assert by_type["Investment"]["migrated_rows"] == 1
    assert by_type["Investment"]["migrated_amount"] == 5000
    # Every remaining difference is explained by a held-back row.
    for row in result.rows:
        if row["Difference"]:
            assert row["Exclusion reasons"]
    assert result.tracker_migrated_rows == 10
    assert result.tracker_manual_rows == 5


def test_reconciliation_attributes_every_exclusion_to_a_reason(
    historical_workbook: Path, legacy_tracker: Path
) -> None:
    run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)

    result = reconcile(historical_workbook, legacy_tracker)

    reasons = " ".join(str(row["Exclusion reasons"]) for row in result.rows)
    assert "formula amount" in reasons
    assert "negative amount read as a credit" in reasons
    for row in result.rows:
        assert row["Original rows"] == row["Migrated rows"] + row["Excluded rows"]
        assert (
            round(float(row["Original amount"]), 2)
            == round(float(row["Migrated amount"]) + float(row["Excluded amount"]), 2)
        )


def test_reconciliation_buckets_held_back_rows(historical_workbook: Path, legacy_tracker: Path) -> None:
    run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)

    buckets = reconcile(historical_workbook, legacy_tracker).buckets

    assert [candidate.description for candidate in buckets[BUCKET_FORMULA]] == ["Magnolia, Osaka"]
    assert [candidate.description for candidate in buckets[BUCKET_INCOME_CREDIT]] == ["Refund"]
    assert [candidate.description for candidate in buckets[BUCKET_RENT_CREDIT]] == ["Sandra"]
    assert [candidate.source_sheet for candidate in buckets[BUCKET_NO_DATE]] == ["Undated"]
    assert [candidate.description for candidate in buckets[BUCKET_SIDEBAR_INCOME]] == [
        "Deposit from school"
    ]


def test_summed_cells_finds_values_inside_a_bounded_sum_range(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "sums.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "United"
    sheet["B1"] = 465.96
    sheet["B2"] = 495
    sheet["C1"] = "=SUM(B1:B2)"
    sheet["E1"] = 1200
    sheet["F1"] = "=SUM(D:D)"
    workbook.save(path)

    covered = summed_cells(load_workbook(path, data_only=False).active)

    assert (1, 2) in covered  # B1, inside the bounded SUM
    assert (2, 2) in covered  # B2
    assert (1, 3) not in covered  # the formula cell itself
    assert (1, 5) not in covered  # E1 is summed by nothing
    # Whole-column ranges such as SUM(D:D) are deliberately ignored: they always point
    # at a transaction block column, which the sidebar scan never looks at anyway.
    assert not any(column == 4 for _, column in covered)


def test_reconciliation_report_files_are_written(
    historical_workbook: Path, legacy_tracker: Path, tmp_path: Path
) -> None:
    run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)
    result = reconcile(historical_workbook, legacy_tracker)

    csv_path = write_reconciliation_csv(result, tmp_path / "recon.csv")
    md_path = write_reconciliation_markdown(result, tmp_path / "recon.md")

    with csv_path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    assert list(rows[0]) == RECONCILIATION_COLUMNS
    assert len(rows) == len(result.rows)

    report = md_path.read_text(encoding="utf-8")
    assert "# Historical reconciliation report" in report
    assert "Raw workbook arithmetic per block" in report
    assert "Totals by Type" in report
    assert BUCKET_FORMULA in report


def test_reconciliation_does_not_modify_either_workbook(
    historical_workbook: Path, legacy_tracker: Path
) -> None:
    run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)
    source_before = historical_workbook.read_bytes()
    tracker_before = legacy_tracker.read_bytes()

    reconcile(historical_workbook, legacy_tracker)

    assert historical_workbook.read_bytes() == source_before
    assert legacy_tracker.read_bytes() == tracker_before


def test_investments_reach_the_tracker_with_period_precision(
    historical_workbook: Path, legacy_tracker: Path
) -> None:
    """Regression guard for the audit fix: dateless Investment blocks must migrate."""
    run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)

    transactions = load_transactions(legacy_tracker)
    migrated = transactions[transactions["Source Sheet"] != ""]
    investments = migrated[migrated["Type"] == "Investment"]
    assert len(investments) == 1
    row = investments.iloc[0]
    assert row["Description"] == "IBKR"
    assert row["Amount"] == 5000
    assert row["Date Precision"] == "Period"
    assert row["Source Sheet"] == "Rent, Income"
    assert row["Source Reference"] == "I3"
    assert str(row["Date"]).startswith("2024-02-01")
