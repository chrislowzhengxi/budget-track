from __future__ import annotations

import csv
from datetime import date
from pathlib import Path
import hashlib

from openpyxl import load_workbook
import pandas as pd
import pytest

from spending_tracker.excel_store import (
    TRANSACTIONS_SHEET,
    load_transactions,
)
from spending_tracker.historical import STATUS_NEEDS_REVIEW
from spending_tracker.migration import (
    REPORT_COLUMNS,
    content_identity,
    dry_run,
    existing_identities,
    provenance_identity,
    run_migration,
)
from spending_tracker.schema import PRECISION_EXACT, TRANSACTION_COLUMNS


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_dry_run_splits_candidates_and_writes_nothing(
    historical_workbook: Path, legacy_tracker: Path
) -> None:
    before_source = _sha256(historical_workbook)
    before_tracker = _sha256(legacy_tracker)

    plan = dry_run(historical_workbook, legacy_tracker)

    assert _sha256(historical_workbook) == before_source
    assert _sha256(legacy_tracker) == before_tracker
    assert len(plan.to_migrate) == 10
    assert {candidate.transaction_type for candidate in plan.to_migrate} == {
        "Expense",
        "Income",
        "Rent",
        "Investment",
    }
    assert len(plan.needs_review) == 5
    assert len(plan.skipped) == 4
    assert plan.duplicates == []


def test_dry_run_summary_reports_counts_and_totals(historical_workbook: Path, legacy_tracker: Path) -> None:
    summary = dry_run(historical_workbook, legacy_tracker).summary

    assert summary["blocks_detected"] == 6
    assert summary["ready_count"] == 10
    assert summary["needs_review_count"] == 5
    assert summary["skipped_count"] == 4
    assert summary["ready_by_type"]["Rent"]["rows"] == 2
    assert summary["ready_by_type"]["Rent"]["amount"] == 1475
    assert summary["ready_by_type"]["Income"]["amount"] == 695.56
    assert summary["ready_by_type"]["Investment"]["rows"] == 1
    assert summary["ready_by_type"]["Investment"]["amount"] == 5000
    assert summary["ready_by_type"]["Expense"]["rows"] == 6
    assert summary["ready_by_period"]["Fall 23"]["rows"] == 4
    assert summary["ready_by_precision"]["Exact"]["rows"] == 5
    assert summary["blank_rows_ignored"] > 0
    assert len(summary["ambiguous_examples"]) == 5


def test_migration_appends_ready_rows_with_provenance(
    historical_workbook: Path, legacy_tracker: Path
) -> None:
    result = run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)

    assert result.migrated == 10
    transactions = load_transactions(legacy_tracker)
    assert len(transactions) == 15  # 5 pre-existing + 10 migrated
    migrated = transactions[transactions["Source Sheet"] != ""]
    assert len(migrated) == 10
    assert set(migrated["Source Period"]) == {"Fall 23", "Winter 24", "Rent, Income"}
    assert list(transactions.columns) == TRANSACTION_COLUMNS
    whole_foods = migrated[migrated["Description"] == "Whole Foods"].iloc[0]
    assert whole_foods["Amount"] == 31.5
    assert whole_foods["Source Reference"] == "C4"
    assert whole_foods["Date Precision"] == PRECISION_EXACT
    assert whole_foods["Added At"]


def test_migration_never_modifies_the_source_workbook(
    historical_workbook: Path, legacy_tracker: Path
) -> None:
    before = _sha256(historical_workbook)
    before_mtime = historical_workbook.stat().st_mtime_ns

    run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)

    assert _sha256(historical_workbook) == before
    assert historical_workbook.stat().st_mtime_ns == before_mtime


def test_migration_refuses_to_write_into_the_source_workbook(historical_workbook: Path) -> None:
    with pytest.raises(ValueError, match="Refusing to migrate into the historical workbook"):
        run_migration(historical_workbook, historical_workbook, apply=True)


def test_migration_backs_up_the_tracker_before_writing(
    historical_workbook: Path, legacy_tracker: Path
) -> None:
    before = _sha256(legacy_tracker)

    result = run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)

    assert result.backup_path is not None
    assert result.backup_path.exists()
    assert _sha256(result.backup_path) == before
    assert result.backup_path.parent.name == "backups"


def test_migration_can_skip_the_backup(historical_workbook: Path, legacy_tracker: Path) -> None:
    result = run_migration(
        historical_workbook,
        legacy_tracker,
        output_dir=legacy_tracker.parent,
        apply=True,
        make_backup=False,
    )

    assert result.backup_path is None
    assert result.migrated == 10


def test_migration_is_idempotent(historical_workbook: Path, legacy_tracker: Path) -> None:
    first = run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)
    after_first = load_transactions(legacy_tracker)

    second = run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)
    after_second = load_transactions(legacy_tracker)

    assert first.migrated == 10
    assert second.migrated == 0
    assert second.duplicates == 10
    assert len(after_second) == len(after_first)
    pd.testing.assert_frame_equal(after_first, after_second)


def test_repeated_purchases_from_different_cells_are_both_migrated(
    historical_workbook: Path, legacy_tracker: Path
) -> None:
    run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)

    transactions = load_transactions(legacy_tracker)
    migrated = transactions[transactions["Source Sheet"] != ""]
    ubers = migrated[migrated["Description"] == "Uber"]
    assert len(ubers) == 2
    assert set(ubers["Source Reference"]) == {"C3", "C6"}


def test_manually_entered_rows_are_not_migrated_twice(
    historical_workbook: Path, legacy_tracker: Path
) -> None:
    # Pretend the user had already typed the Winter 24 Tatte row by hand.
    from spending_tracker.excel_store import append_transactions

    append_transactions(
        legacy_tracker,
        [
            {
                "Date": date(2024, 1, 10),
                "Description": "Tatte",
                "Amount": 8.5,
                "Type": "Expense",
                "Source Line": "tatte 8.50",
            }
        ],
    )

    result = run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)

    transactions = load_transactions(legacy_tracker)
    assert result.migrated == 9
    assert len(transactions[transactions["Description"] == "Tatte"]) == 2  # legacy row + the manual one


def test_needs_review_rows_are_written_to_the_review_csv(
    historical_workbook: Path, legacy_tracker: Path
) -> None:
    result = run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)

    with result.review_path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))

    assert list(rows[0]) == REPORT_COLUMNS
    assert len(rows) == 9  # 5 needing review + 4 skipped
    review_statuses = {row["Status"] for row in rows}
    assert review_statuses == {STATUS_NEEDS_REVIEW, "Skipped"}
    undated = next(row for row in rows if row["Source Sheet"] == "Undated")
    assert undated["Description"] == "IBKR"
    assert undated["Date"] == ""
    assert undated["Date Precision"] == "Unknown"

    migrated_descriptions = set(load_transactions(legacy_tracker)["Description"])
    assert "Refund" not in migrated_descriptions  # negative amounts stay in review
    assert "Deposit from school" not in migrated_descriptions  # so do sidebar values


def test_report_artifacts_cover_every_candidate(historical_workbook: Path, legacy_tracker: Path) -> None:
    result = run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)

    with result.report_csv_path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 10 + 5 + 4
    assert sum(1 for row in rows if row["Status"] == "Migrated") == 10

    report = result.report_md_path.read_text(encoding="utf-8")
    assert "# Historical migration report" in report
    assert "Ready (auto-migrated): 10" in report
    assert "Date handling" in report
    assert str(historical_workbook) in report


def test_second_run_report_marks_rows_as_duplicates(
    historical_workbook: Path, legacy_tracker: Path
) -> None:
    run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)
    result = run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)

    with result.report_csv_path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    assert sum(1 for row in rows if row["Status"].startswith("Duplicate")) == 10
    assert sum(1 for row in rows if row["Status"] == "Migrated") == 0


def test_migration_creates_the_transactions_sheet_when_missing(
    historical_workbook: Path, empty_tracker: Path
) -> None:
    result = run_migration(historical_workbook, empty_tracker, output_dir=empty_tracker.parent, apply=True)

    workbook = load_workbook(empty_tracker)
    assert TRANSACTIONS_SHEET in workbook.sheetnames
    assert workbook["24-25 Fall, Winter"]["B1"].value == "Fall 24"
    assert result.migrated == 10


def test_identity_helpers() -> None:
    assert provenance_identity("25-26", "C4") == ("25-26", "C4")
    assert provenance_identity("", "C4") is None
    assert provenance_identity("25-26", None) is None
    assert content_identity(date(2026, 1, 2), " Whole  Foods ", 12.004, "Expense") == (
        "2026-01-02",
        "whole foods",
        12.0,
        "Expense",
    )
    assert content_identity(pd.Timestamp("2026-01-02"), "Uber", 12, "Expense")[0] == "2026-01-02"
    assert content_identity(None, "Uber", None, "Expense") == ("", "uber", 0.0, "Expense")


def test_existing_identities_uses_content_only_for_hand_entered_rows() -> None:
    frame = pd.DataFrame(
        [
            {
                "Date": date(2026, 1, 1),
                "Description": "Uber",
                "Amount": 12,
                "Type": "Expense",
                "Source Sheet": "25-26",
                "Source Reference": "C4",
            },
            {
                "Date": date(2026, 1, 2),
                "Description": "Tatte",
                "Amount": 8.5,
                "Type": "Expense",
                "Source Sheet": "",
                "Source Reference": "",
            },
        ]
    )

    provenance, content = existing_identities(frame)

    assert provenance == {("25-26", "C4")}
    assert content == {("2026-01-02", "tatte", 8.5, "Expense")}


def test_dry_run_reports_duplicates_without_touching_the_tracker(
    historical_workbook: Path, legacy_tracker: Path
) -> None:
    run_migration(historical_workbook, legacy_tracker, output_dir=legacy_tracker.parent, apply=True)
    before = _sha256(legacy_tracker)

    plan = dry_run(historical_workbook, legacy_tracker)

    assert len(plan.duplicates) == 10
    assert plan.to_migrate == []
    assert _sha256(legacy_tracker) == before


def test_applying_archives_a_timestamped_report(historical_workbook: Path, legacy_tracker: Path) -> None:
    directory = legacy_tracker.parent

    run_migration(historical_workbook, legacy_tracker, output_dir=directory, apply=True)

    archives = list(directory.glob("historical_migration_report-*.md"))
    assert len(archives) == 1
    assert "Ready (auto-migrated): 10" in archives[0].read_text(encoding="utf-8")

    # A re-run migrates nothing, so it does not add another archive.
    run_migration(historical_workbook, legacy_tracker, output_dir=directory, apply=True)
    assert len(list(directory.glob("historical_migration_report-*.md"))) == 1
