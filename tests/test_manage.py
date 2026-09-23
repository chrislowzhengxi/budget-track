from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pandas as pd
import pytest

from spending_tracker.excel_store import append_transactions, load_transactions
from spending_tracker.manage import (
    audit_log_path,
    delete_transactions,
    read_audit,
    recategorize,
    record_audit,
    save_transactions,
    search_transactions,
    update_transaction,
)


def _tracker(tmp_path: Path) -> Path:
    path = tmp_path / "tracker.xlsx"
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = "History"
    workbook.active["A1"] = "keep me"
    workbook.save(path)
    append_transactions(
        path,
        [
            {
                "Date": date(2026, 9, 1),
                "Description": "Uber",
                "Amount": 12.0,
                "Type": "Expense",
                "Category": "Transportation",
                "Source Line": "uber 12",
            },
            {
                "Date": date(2026, 9, 2),
                "Description": "Tatte",
                "Amount": 8.5,
                "Type": "Expense",
                "Category": "Other",
                "Source Line": "tatte 8.50",
            },
            {
                "Date": date(2026, 9, 3),
                "Description": "TA",
                "Amount": 695.56,
                "Type": "Income",
                "Source Line": "TA (695.56)",
            },
            {
                "Date": date(2026, 10, 1),
                "Description": "Whole Foods",
                "Amount": 31.0,
                "Type": "Expense",
                "Category": "Groceries",
                "Source Line": "whole foods 31",
            },
        ],
    )
    return path


def test_search_by_description_type_category_and_dates(tmp_path: Path) -> None:
    frame = load_transactions(_tracker(tmp_path))

    assert search_transactions(frame, text="tat")["Description"].tolist() == ["Tatte"]
    assert search_transactions(frame, text="UBER")["Description"].tolist() == ["Uber"]
    assert search_transactions(frame, types=["Income"])["Description"].tolist() == ["TA"]
    assert search_transactions(frame, categories=["Groceries"])["Description"].tolist() == ["Whole Foods"]
    assert search_transactions(frame, start=date(2026, 9, 2), end=date(2026, 9, 3))[
        "Description"
    ].tolist() == ["Tatte", "TA"]
    assert len(search_transactions(frame)) == 4
    assert search_transactions(pd.DataFrame()).empty


def test_search_matches_the_source_line_too(tmp_path: Path) -> None:
    frame = load_transactions(_tracker(tmp_path))

    assert search_transactions(frame, text="695.56")["Description"].tolist() == ["TA"]


def test_update_transaction_changes_only_the_requested_row(tmp_path: Path) -> None:
    frame = load_transactions(_tracker(tmp_path))

    updated = update_transaction(frame, 1, {"Category": "Dining", "Amount": 9.25})

    assert updated.loc[1, "Category"] == "Dining"
    assert updated.loc[1, "Amount"] == 9.25
    assert updated.loc[0, "Category"] == "Transportation"
    assert frame.loc[1, "Amount"] == 8.5  # original frame untouched


def test_update_transaction_rejects_bad_input(tmp_path: Path) -> None:
    frame = load_transactions(_tracker(tmp_path))

    with pytest.raises(KeyError):
        update_transaction(frame, 99, {"Category": "Dining"})
    with pytest.raises(ValueError, match="not editable"):
        update_transaction(frame, 0, {"Source Sheet": "hacked"})
    with pytest.raises(ValueError, match="greater than zero"):
        update_transaction(frame, 0, {"Amount": 0})


def test_delete_transactions_removes_rows_and_reindexes(tmp_path: Path) -> None:
    frame = load_transactions(_tracker(tmp_path))

    remaining = delete_transactions(frame, [1])

    assert remaining["Description"].tolist() == ["Uber", "TA", "Whole Foods"]
    assert remaining.index.tolist() == [0, 1, 2]
    with pytest.raises(KeyError):
        delete_transactions(frame, [42])


def test_save_transactions_backs_up_writes_and_audits(tmp_path: Path) -> None:
    path = _tracker(tmp_path)
    frame = load_transactions(path)

    written = save_transactions(path, delete_transactions(frame, [0]), action="delete", details={"rows": 1})

    reloaded = load_transactions(path)
    assert written == 3
    assert "Uber" not in set(reloaded["Description"])
    assert list((tmp_path / "backups").glob("*.xlsx"))
    entries = read_audit(path)
    assert entries[0]["action"] == "delete"
    assert entries[0]["details"]["rows"] == 1
    assert entries[0]["details"]["rows_after"] == 3
    assert "backup" in entries[0]["details"]
    from openpyxl import load_workbook

    assert load_workbook(path)["History"]["A1"].value == "keep me"


def test_save_transactions_can_skip_the_backup(tmp_path: Path) -> None:
    path = _tracker(tmp_path)

    save_transactions(path, load_transactions(path), action="noop", make_backup=False)

    assert not (tmp_path / "backups").exists()


def test_record_and_read_audit_return_newest_first(tmp_path: Path) -> None:
    path = tmp_path / "tracker.xlsx"

    record_audit(path, "first", {"n": 1}, now=datetime(2026, 9, 1, 10, 0, 0))
    record_audit(path, "second", {"when": date(2026, 9, 2)}, now=datetime(2026, 9, 2, 10, 0, 0))

    entries = read_audit(path)
    assert [entry["action"] for entry in entries] == ["second", "first"]
    assert entries[0]["details"]["when"] == "2026-09-02"
    assert entries[0]["at"] == "2026-09-02 10:00:00"
    assert read_audit(path, limit=1) == entries[:1]
    assert audit_log_path(path).exists()
    # The log lives beside the workbook, so an untouched folder has no history.
    other = tmp_path / "elsewhere" / "tracker.xlsx"
    other.parent.mkdir()
    assert read_audit(other) == []


def test_recategorize_fills_blanks_without_touching_manual_values(tmp_path: Path) -> None:
    path = _tracker(tmp_path)

    result = recategorize(path)

    frame = load_transactions(path)
    assert result["changed"] == 0  # every expense already had a category
    assert frame.loc[1, "Category"] == "Other"

    overwritten = recategorize(path, overwrite=True)
    frame = load_transactions(path)
    assert overwritten["changed"] == 1
    assert frame.loc[1, "Category"] == "Dining"
    assert read_audit(path)[0]["action"] == "recategorize"


def test_recategorize_applies_user_rules(tmp_path: Path) -> None:
    path = _tracker(tmp_path)

    recategorize(path, overwrite=True, user_rules={"tatte": "Groceries"})

    assert load_transactions(path).loc[1, "Category"] == "Groceries"


def test_recategorize_on_an_empty_workbook_is_a_no_op(tmp_path: Path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "blank.xlsx"
    Workbook().save(path)

    assert recategorize(path) == {"rows": 0, "changed": 0, "written": 0}
